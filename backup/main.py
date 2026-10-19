import os
import io
import sqlite3
import hashlib
from datetime import date, datetime
from typing import Optional, List
from fastapi import FastAPI, HTTPException
from fastapi.responses import HTMLResponse, StreamingResponse
from pydantic import BaseModel
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

app = FastAPI(title="Presales Activities Hub")

if os.path.exists("/home") and os.path.isdir("/home"):
    DB_FILE = "/home/presales_tracker.db"
else:
    DB_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "presales_tracker.db")

def hash_pw(pw: str) -> str:
    return hashlib.sha256(pw.encode()).hexdigest()

def init_db():
    with sqlite3.connect(DB_FILE) as conn:
        conn.execute("PRAGMA foreign_keys = ON;")
        cur = conn.cursor()
        cur.execute("""
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT UNIQUE NOT NULL,
            password_hash TEXT NOT NULL,
            full_name TEXT NOT NULL,
            role TEXT CHECK(role IN ('admin', 'view_only')) NOT NULL DEFAULT 'view_only',
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )""")
        cur.execute("""
        CREATE TABLE IF NOT EXISTS customers (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            industry TEXT,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )""")
        cur.execute("""
        CREATE TABLE IF NOT EXISTS opportunities (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            customer_id INTEGER NOT NULL,
            title TEXT NOT NULL,
            status TEXT DEFAULT 'New' CHECK(status IN ('New', 'On progress', 'Hold', 'Completed')),
            priority TEXT DEFAULT 'Medium' CHECK(priority IN ('Urgent', 'High', 'Medium', 'Low')),
            strategy_notes TEXT DEFAULT '',
            end_date DATE,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (customer_id) REFERENCES customers(id) ON DELETE CASCADE
        )""")
        cur.execute("""
        CREATE TABLE IF NOT EXISTS stage_tasks (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            opportunity_id INTEGER NOT NULL,
            stage_id TEXT NOT NULL CHECK(stage_id IN ('Stage 02 Assessment', 'Stage 03 POC', 'Stage 04 Designing', 'Stage 05 Proposal')),
            title TEXT NOT NULL,
            start_date DATE,
            end_date DATE,
            status TEXT DEFAULT 'New' CHECK(status IN ('New', 'On going', 'Hold', 'Completed')),
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (opportunity_id) REFERENCES opportunities(id) ON DELETE CASCADE
        )""")
        cur.execute("""
        CREATE TABLE IF NOT EXISTS task_daily_updates (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            task_id INTEGER NOT NULL,
            update_date DATE DEFAULT (CURRENT_DATE),
            notes TEXT NOT NULL,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (task_id) REFERENCES stage_tasks(id) ON DELETE CASCADE
        )""")
        cur.execute("""
        CREATE TABLE IF NOT EXISTS activity_audit_logs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            customer_name TEXT,
            opp_title TEXT,
            action_type TEXT NOT NULL,
            details TEXT NOT NULL,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )""")
        conn.commit()

        cur.execute("SELECT id FROM users WHERE username = 'admin'")
        if not cur.fetchone():
            cur.execute("""
                INSERT INTO users (username, password_hash, full_name, role)
                VALUES ('admin', ?, 'Chief Administrator', 'admin')
            """, (hash_pw('admin123'),))
            conn.commit()

        cur.execute("PRAGMA table_info(opportunities);")
        columns = [row[1] for row in cur.fetchall()]
        if "end_date" not in columns:
            cur.execute("ALTER TABLE opportunities ADD COLUMN end_date DATE;")
        conn.commit()

init_db()

def log_audit(action: str, details: str, cust_name: str = "", opp_title: str = ""):
    try:
        with sqlite3.connect(DB_FILE) as conn:
            cur = conn.cursor()
            cur.execute("""
                INSERT INTO activity_audit_logs (customer_name, opp_title, action_type, details)
                VALUES (?, ?, ?, ?)
            """, (cust_name, opp_title, action, details))
            conn.commit()
    except Exception as e:
        print(f"Audit log error: {e}")

# --- Pydantic Data Models ---
class LoginRequest(BaseModel):
    username: str
    password: str

class UserCreate(BaseModel):
    username: str
    password: str
    full_name: str
    role: str

class UnifiedEngagementCreate(BaseModel):
    is_new_customer: bool
    customer_id: Optional[int] = None
    customer_name: Optional[str] = None
    industry: Optional[str] = "General"
    opp_title: str
    priority: Optional[str] = "Medium"
    status: Optional[str] = "New"
    end_date: Optional[str] = None

class CustomerUpdate(BaseModel):
    name: str
    industry: Optional[str] = "General"

class OpportunityUpdate(BaseModel):
    title: Optional[str] = None
    status: Optional[str] = None
    priority: Optional[str] = None
    strategy_notes: Optional[str] = None
    end_date: Optional[str] = None

class TaskCreate(BaseModel):
    opportunity_id: int
    stage_id: str
    title: str
    start_date: Optional[str] = None
    end_date: Optional[str] = None
    status: Optional[str] = "New"

class TaskEdit(BaseModel):
    title: str
    start_date: Optional[str] = None
    end_date: Optional[str] = None
    status: str

class CommentCreate(BaseModel):
    notes: str

# --- Authentication & Users ---
@app.post("/api/auth/login")
def login(req: LoginRequest):
    with sqlite3.connect(DB_FILE) as conn:
        conn.row_factory = sqlite3.Row
        cur = conn.cursor()
        cur.execute("SELECT id, username, full_name, role FROM users WHERE username = ? AND password_hash = ?", 
                    (req.username.strip(), hash_pw(req.password.strip())))
        user = cur.fetchone()
        if not user:
            raise HTTPException(status_code=401, detail="Invalid username or password")
        return dict(user)

@app.get("/api/users")
def get_users():
    with sqlite3.connect(DB_FILE) as conn:
        conn.row_factory = sqlite3.Row
        cur = conn.cursor()
        cur.execute("SELECT id, username, full_name, role, created_at FROM users ORDER BY id ASC")
        return [dict(r) for r in cur.fetchall()]

@app.post("/api/users")
def create_user(u: UserCreate):
    with sqlite3.connect(DB_FILE) as conn:
        cur = conn.cursor()
        try:
            cur.execute("""
                INSERT INTO users (username, password_hash, full_name, role)
                VALUES (?, ?, ?, ?)
            """, (u.username.strip().lower(), hash_pw(u.password.strip()), u.full_name.strip(), u.role))
            conn.commit()
            uid = cur.lastrowid
        except sqlite3.IntegrityError:
            raise HTTPException(status_code=400, detail="Username already exists")
    log_audit("USER_CREATED", f"User account '{u.username}' created with role '{u.role}'")
    return {"id": uid, "username": u.username, "role": u.role}

# --- State Endpoint ---
@app.get("/api/state")
def get_full_state():
    with sqlite3.connect(DB_FILE) as conn:
        conn.row_factory = sqlite3.Row
        cur = conn.cursor()
        cur.execute("SELECT * FROM customers ORDER BY name ASC")
        customers = [dict(r) for r in cur.fetchall()]

        cur.execute("SELECT * FROM opportunities ORDER BY id DESC")
        opps = [dict(r) for r in cur.fetchall()]

        cur.execute("SELECT * FROM stage_tasks ORDER BY id ASC")
        tasks = [dict(r) for r in cur.fetchall()]

        cur.execute("SELECT * FROM task_daily_updates ORDER BY id DESC")
        updates = [dict(r) for r in cur.fetchall()]

        cur.execute("SELECT * FROM activity_audit_logs ORDER BY id DESC LIMIT 50")
        logs = [dict(r) for r in cur.fetchall()]

        task_updates_map = {}
        for u in updates:
            task_updates_map.setdefault(u["task_id"], []).append({
                "id": u["id"],
                "date": u["update_date"],
                "text": u["notes"]
            })

        for t in tasks:
            t["comments"] = task_updates_map.get(t["id"], [])

        return {
            "customers": customers,
            "opportunities": opps,
            "tasks": tasks,
            "audit_logs": logs
        }

@app.post("/api/engagements/unified")
def create_unified_engagement(payload: UnifiedEngagementCreate):
    with sqlite3.connect(DB_FILE) as conn:
        cur = conn.cursor()
        if payload.is_new_customer:
            cust_name = payload.customer_name.strip()
            cur.execute("INSERT INTO customers (name, industry) VALUES (?, ?)", (cust_name, payload.industry.strip()))
            cust_id = cur.lastrowid
            log_audit("CUSTOMER_CREATED", f"Account '{cust_name}' created", cust_name=cust_name)
        else:
            cust_id = payload.customer_id
            cur.execute("SELECT name FROM customers WHERE id = ?", (cust_id,))
            row = cur.fetchone()
            if not row:
                raise HTTPException(status_code=404, detail="Customer not found")
            cust_name = row[0]

        cur.execute("""
            INSERT INTO opportunities (customer_id, title, status, priority, end_date)
            VALUES (?, ?, ?, ?, ?)
        """, (cust_id, payload.opp_title.strip(), payload.status, payload.priority, payload.end_date))
        conn.commit()
        opp_id = cur.lastrowid

    log_audit("OPP_CREATED", f"Opportunity '{payload.opp_title}' added", cust_name=cust_name, opp_title=payload.opp_title)
    return {"customer_id": cust_id, "opportunity_id": opp_id}

@app.put("/api/customers/{cust_id}")
def update_customer(cust_id: int, c: CustomerUpdate):
    with sqlite3.connect(DB_FILE) as conn:
        conn.row_factory = sqlite3.Row
        cur = conn.cursor()
        cur.execute("SELECT name FROM customers WHERE id = ?", (cust_id,))
        old_cust = cur.fetchone()
        if not old_cust:
            raise HTTPException(status_code=404, detail="Customer not found")
        
        cur.execute("UPDATE customers SET name = ?, industry = ? WHERE id = ?", (c.name.strip(), c.industry.strip(), cust_id))
        conn.commit()

    log_audit("CUSTOMER_UPDATED", f"Customer details updated from '{old_cust['name']}' to '{c.name}'", cust_name=c.name)
    return {"status": "success"}

@app.patch("/api/opportunities/{opp_id}")
def update_opportunity(opp_id: int, payload: OpportunityUpdate):
    with sqlite3.connect(DB_FILE) as conn:
        conn.row_factory = sqlite3.Row
        cur = conn.cursor()
        cur.execute("SELECT o.*, c.name as customer_name FROM opportunities o JOIN customers c ON o.customer_id = c.id WHERE o.id = ?", (opp_id,))
        opp = cur.fetchone()
        if not opp:
            raise HTTPException(status_code=404, detail="Opportunity not found")
        
        updates = []
        params = []
        if payload.title is not None:
            updates.append("title = ?")
            params.append(payload.title.strip())
        if payload.status is not None:
            updates.append("status = ?")
            params.append(payload.status)
        if payload.priority is not None:
            updates.append("priority = ?")
            params.append(payload.priority)
        if payload.strategy_notes is not None:
            updates.append("strategy_notes = ?")
            params.append(payload.strategy_notes)
        if payload.end_date is not None:
            updates.append("end_date = ?")
            params.append(payload.end_date)
        
        if updates:
            params.append(opp_id)
            cur.execute(f"UPDATE opportunities SET {', '.join(updates)} WHERE id = ?", params)
            conn.commit()
    
    log_audit("OPP_UPDATED", f"Opportunity updated: {payload.dict(exclude_unset=True)}", cust_name=opp["customer_name"], opp_title=opp["title"])
    return {"status": "success"}

@app.post("/api/tasks")
def create_task(t: TaskCreate):
    with sqlite3.connect(DB_FILE) as conn:
        conn.row_factory = sqlite3.Row
        cur = conn.cursor()
        cur.execute("SELECT o.title, c.name as customer_name FROM opportunities o JOIN customers c ON o.customer_id = c.id WHERE o.id = ?", (t.opportunity_id,))
        info = cur.fetchone()
        cur.execute("""
            INSERT INTO stage_tasks (opportunity_id, stage_id, title, start_date, end_date, status)
            VALUES (?, ?, ?, ?, ?, ?)
        """, (t.opportunity_id, t.stage_id, t.title.strip(), t.start_date, t.end_date, t.status))
        conn.commit()
        task_id = cur.lastrowid
    
    log_audit("TASK_CREATED", f"Task '{t.title}' created in {t.stage_id}", cust_name=info["customer_name"] if info else "", opp_title=info["title"] if info else "")
    return {"id": task_id}

@app.put("/api/tasks/{task_id}")
def edit_task(task_id: int, payload: TaskEdit):
    with sqlite3.connect(DB_FILE) as conn:
        conn.row_factory = sqlite3.Row
        cur = conn.cursor()
        cur.execute("""
            SELECT t.title, t.stage_id, o.title as opp_title, c.name as customer_name 
            FROM stage_tasks t 
            JOIN opportunities o ON t.opportunity_id = o.id 
            JOIN customers c ON o.customer_id = c.id 
            WHERE t.id = ?
        """, (task_id,))
        task = cur.fetchone()
        if not task:
            raise HTTPException(status_code=404, detail="Task not found")
        
        cur.execute("""
            UPDATE stage_tasks 
            SET title = ?, start_date = ?, end_date = ?, status = ?
            WHERE id = ?
        """, (payload.title.strip(), payload.start_date, payload.end_date, payload.status, task_id))
        conn.commit()

    log_audit("TASK_UPDATED", f"Task '{payload.title}' modified", cust_name=task["customer_name"], opp_title=task["opp_title"])
    return {"status": "success"}

@app.delete("/api/tasks/{task_id}")
def delete_task(task_id: int):
    with sqlite3.connect(DB_FILE) as conn:
        conn.row_factory = sqlite3.Row
        cur = conn.cursor()
        cur.execute("""
            SELECT t.title, o.title as opp_title, c.name as customer_name 
            FROM stage_tasks t 
            JOIN opportunities o ON t.opportunity_id = o.id 
            JOIN customers c ON o.customer_id = c.id 
            WHERE t.id = ?
        """, (task_id,))
        task = cur.fetchone()
        if task:
            cur.execute("DELETE FROM stage_tasks WHERE id = ?", (task_id,))
            conn.commit()
            log_audit("TASK_DELETED", f"Task '{task['title']}' removed", cust_name=task["customer_name"], opp_title=task["opp_title"])
    return {"status": "deleted"}

@app.post("/api/tasks/{task_id}/comments")
def add_task_comment(task_id: int, comment: CommentCreate):
    with sqlite3.connect(DB_FILE) as conn:
        conn.row_factory = sqlite3.Row
        cur = conn.cursor()
        cur.execute("""
            SELECT t.title, o.title as opp_title, c.name as customer_name 
            FROM stage_tasks t 
            JOIN opportunities o ON t.opportunity_id = o.id 
            JOIN customers c ON o.customer_id = c.id 
            WHERE t.id = ?
        """, (task_id,))
        task = cur.fetchone()
        cur.execute("""
            INSERT INTO task_daily_updates (task_id, notes) VALUES (?, ?)
        """, (task_id, comment.notes.strip()))
        conn.commit()
    
    log_audit("DAILY_UPDATE_POSTED", f"Update logged for task '{task['title']}': {comment.notes.strip()}",
              cust_name=task["customer_name"] if task else "", opp_title=task["opp_title"] if task else "")
    return {"status": "success"}

# --- Executive Exports ---
def fetch_executive_report_data():
    with sqlite3.connect(DB_FILE) as conn:
        conn.row_factory = sqlite3.Row
        cur = conn.cursor()
        cur.execute("SELECT COUNT(DISTINCT customer_id) as total_cust FROM opportunities WHERE status != 'Completed'")
        active_customers = cur.fetchone()["total_cust"]
        cur.execute("SELECT COUNT(*) as total_opps FROM opportunities WHERE status != 'Completed'")
        active_opps = cur.fetchone()["total_opps"]
        cur.execute("SELECT COUNT(*) as urgent_deals FROM opportunities WHERE priority IN ('Urgent', 'High') AND status != 'Completed'")
        urgent_deals = cur.fetchone()["urgent_deals"]
        cur.execute("SELECT stage_id, COUNT(*) as count FROM stage_tasks GROUP BY stage_id")
        stage_counts = {row["stage_id"]: row["count"] for row in cur.fetchall()}
        cur.execute("""
            SELECT c.name as customer_name, o.title as opp_title, o.status, o.priority, o.end_date
            FROM opportunities o
            JOIN customers c ON o.customer_id = c.id
            ORDER BY c.name ASC
        """)
        opps = [dict(r) for r in cur.fetchall()]
        return {
            "date": str(date.today()),
            "active_customers": active_customers,
            "active_opps": active_opps,
            "urgent_deals": urgent_deals,
            "stage_counts": stage_counts,
            "opportunities": opps
        }

@app.get("/api/reports/export/pdf")
def export_executive_pdf():
    data = fetch_executive_report_data()
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(letter), leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('DocTitle', parent=styles['Heading1'], fontSize=18, textColor=colors.HexColor('#0f172a'))
    
    story = [
        Paragraph("Presales Activities Hub - Operations Briefing", title_style),
        Paragraph(f"Generated on {data['date']} | Executive Operations Summary", styles['Normal']),
        Spacer(1, 15)
    ]
    summary_data = [
        ["Active Accounts", "Pipeline Opportunities", "Urgent / High Attention", "Total Stage Tasks"],
        [str(data['active_customers']), str(data['active_opps']), str(data['urgent_deals']), sum(data['stage_counts'].values())]
    ]
    summary_table = Table(summary_data, colWidths=[180, 180, 180, 180])
    summary_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#fff7ed')),
        ('TEXTCOLOR', (0, 0), (-1, -1), colors.HexColor('#0f172a')),
        ('FONTNAME', (0, 0), (-1, -1), 'Helvetica-Bold'),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#fdba74')),
    ]))
    story.append(summary_table)
    story.append(Spacer(1, 20))

    opp_rows = [["Customer Account", "Opportunity Title", "Target End Date", "Status", "Priority"]]
    for opp in data["opportunities"]:
        opp_rows.append([opp["customer_name"], opp["opp_title"], opp['end_date'] or 'Not Set', opp["status"], opp["priority"]])
    opp_table = Table(opp_rows, colWidths=[180, 260, 140, 80, 80])
    opp_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0f172a')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, -1), 'Helvetica-Bold'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
        ('FONTSIZE', (0, 1), (-1, -1), 8),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')])
    ]))
    story.append(opp_table)
    doc.build(story)
    buffer.seek(0)
    return StreamingResponse(buffer, media_type="application/pdf", headers={"Content-Disposition": f"attachment; filename=Presales_Hub_{data['date']}.pdf"})

@app.get("/api/reports/export/pptx")
def export_executive_pptx():
    data = fetch_executive_report_data()
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.8))
    p = tb.text_frame.paragraphs[0]
    p.text = f"Presales Activities Hub - Executive Briefing ({data['date']})"
    p.font.size = Pt(22)
    p.font.bold = True

    metrics = [("Active Accounts", str(data["active_customers"])), ("Pipeline Deals", str(data["active_opps"])), ("Urgent Deals", str(data["urgent_deals"]))]
    for i, (lbl, val) in enumerate(metrics):
        box = slide.shapes.add_textbox(Inches(0.8 + (i * 4.0)), Inches(1.5), Inches(3.6), Inches(1.2))
        tf = box.text_frame
        p_val = tf.paragraphs[0]
        p_val.text = val
        p_val.font.size = Pt(36)
        p_val.font.bold = True
        p_lbl = tf.add_paragraph()
        p_lbl.text = lbl
        p_lbl.font.size = Pt(11)

    rows = min(len(data["opportunities"]) + 1, 9)
    table_shape = slide.shapes.add_table(rows, 5, Inches(0.8), Inches(3.0), Inches(11.7), Inches(3.8))
    table = table_shape.table
    for c, h in enumerate(["Customer", "Opportunity", "Target End Date", "Status", "Priority"]):
        table.cell(0, c).text = h
        table.cell(0, c).fill.solid()
        table.cell(0, c).fill.fore_color.rgb = RGBColor(254, 215, 170)
    for r_idx, opp in enumerate(data["opportunities"][:rows - 1]):
        table.cell(r_idx + 1, 0).text = opp["customer_name"]
        table.cell(r_idx + 1, 1).text = opp["opp_title"]
        table.cell(r_idx + 1, 2).text = opp['end_date'] or '-'
        table.cell(r_idx + 1, 3).text = opp["status"]
        table.cell(r_idx + 1, 4).text = opp["priority"]
        for col in range(5):
            table.cell(r_idx + 1, col).text_frame.paragraphs[0].font.size = Pt(10)

    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation", headers={"Content-Disposition": f"attachment; filename=Presales_Hub_{data['date']}.pptx"})

# --- Frontend with Kaleido-Style Fluid Mesh Login ---
@app.get("/", response_class=HTMLResponse)
def index_page():
    return """
    <!DOCTYPE html>
    <html lang="en">
    <head>
      <meta charset="UTF-8">
      <title>Presales Activities Hub</title>
      <link rel="icon" type="image/svg+xml" href="data:image/svg+xml,<svg xmlns='http://www.w3.org/2000/svg' viewBox='0 0 100 100'><defs><linearGradient id='g' x1='0%' y1='0%' x2='100%' y2='100%'><stop offset='0%' stop-color='%23fb923c'/><stop offset='100%' stop-color='%234f46e5'/></linearGradient></defs><rect width='100' height='100' rx='25' fill='url(%23g)'/><path d='M30 70 L50 25 L70 70 L50 55 Z' fill='white'/></svg>">
      
      <link rel="preconnect" href="https://fonts.googleapis.com">
      <link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
      <link href="https://fonts.googleapis.com/css2?family=Plus+Jakarta+Sans:wght@400;500;600;700;800;900&display=swap" rel="stylesheet">

      <script src="https://cdn.tailwindcss.com"></script>
      <link rel="stylesheet" href="https://cdnjs.cloudflare.com/ajax/libs/font-awesome/6.4.0/css/all.min.css">
      
      <style>
        body {
          font-family: 'Plus Jakarta Sans', sans-serif;
          background-color: #f8fafc;
          background-image: 
            radial-gradient(at 0% 0%, rgba(251, 146, 60, 0.06) 0px, transparent 50%),
            radial-gradient(at 100% 0%, rgba(99, 102, 241, 0.06) 0px, transparent 50%),
            linear-gradient(to right, rgba(226, 232, 240, 0.6) 1px, transparent 1px),
            linear-gradient(to bottom, rgba(226, 232, 240, 0.6) 1px, transparent 1px);
          background-size: 100% 100%, 100% 100%, 32px 32px, 32px 32px;
        }

        .custom-scrollbar::-webkit-scrollbar { width: 5px; height: 5px; }
        .custom-scrollbar::-webkit-scrollbar-track { background: #f1f5f9; border-radius: 8px; }
        .custom-scrollbar::-webkit-scrollbar-thumb { background: #fdba74; border-radius: 8px; }
        .custom-scrollbar::-webkit-scrollbar-thumb:hover { background: #fb923c; }

        .glass-card {
          background: rgba(255, 255, 255, 0.94);
          backdrop-filter: blur(14px);
          -webkit-backdrop-filter: blur(14px);
          border: 1px solid rgba(226, 232, 240, 0.95);
        }

        .tile-card {
          box-shadow: 0 10px 25px -5px rgba(15, 23, 42, 0.05), 0 8px 10px -6px rgba(15, 23, 42, 0.03);
        }
        .tile-card:hover {
          box-shadow: 0 20px 30px -10px rgba(15, 23, 42, 0.1), 0 10px 15px -5px rgba(251, 146, 60, 0.12);
        }

        /* Vivid Fluid Mesh Canvas matching user image */
        .fluid-ink-bg {
          background-color: #0c1427;
          background-image: 
            radial-gradient(circle at 15% 45%, rgba(249, 115, 22, 0.8) 0%, transparent 45%),
            radial-gradient(circle at 85% 65%, rgba(236, 72, 153, 0.75) 0%, transparent 50%),
            radial-gradient(circle at 50% 30%, rgba(14, 165, 233, 0.8) 0%, transparent 50%),
            radial-gradient(circle at 75% 15%, rgba(99, 102, 241, 0.85) 0%, transparent 45%),
            radial-gradient(circle at 35% 85%, rgba(239, 68, 68, 0.75) 0%, transparent 50%),
            url('https://images.unsplash.com/photo-1541701494587-cb58502866ab?auto=format&fit=crop&w=1920&q=85');
          background-size: cover;
          background-position: center;
          background-blend-mode: color-dodge, screen, hard-light, normal;
        }

        .kaleido-glass {
          background: rgba(255, 255, 255, 0.08);
          backdrop-filter: blur(30px);
          -webkit-backdrop-filter: blur(30px);
          border: 1px solid rgba(255, 255, 255, 0.18);
          box-shadow: 0 30px 60px -12px rgba(0, 0, 0, 0.45);
        }

        * { 
          transition-property: background-color, border-color, color, fill, stroke, opacity, box-shadow, transform; 
          transition-duration: 180ms; 
          transition-timing-function: cubic-bezier(0.4, 0, 0.2, 1); 
        }
      </style>
    </head>
    <body class="text-slate-800 min-h-screen selection:bg-orange-400 selection:text-white">

      <!-- ================= KALEIDO-STYLE MODERN VIVID LOGIN SCREEN ================= -->
      <div id="loginOverlay" class="fixed inset-0 z-50 flex items-center justify-center p-6 md:p-12 fluid-ink-bg overflow-y-auto">
        
        <div class="w-full max-w-6xl flex flex-col lg:flex-row items-center justify-between gap-12 my-auto">
          
          <!-- Left Hero Section (matching user image layout) -->
          <div class="w-full lg:w-1/2 space-y-8 text-white">
            <!-- Brand Badge -->
            <div class="flex items-center gap-2.5">
              <div class="w-7 h-7 rounded-lg bg-gradient-to-tr from-amber-400 to-orange-500 flex items-center justify-center shadow-lg shadow-orange-500/30">
                <i class="fa-solid fa-shapes text-white text-xs"></i>
              </div>
              <span class="font-extrabold tracking-wide text-base">Presales Hub</span>
            </div>

            <!-- Overline Tag -->
            <div class="space-y-3">
              <span class="text-[11px] font-black uppercase tracking-widest text-slate-300/80">LIFECYCLE INTELLIGENCE FOR ENGINEERING TEAMS</span>
              <h1 class="text-4xl md:text-5xl lg:text-6xl font-black leading-[1.1] tracking-tight text-white drop-shadow-md">
                Every opportunity your team touches, in one place.
              </h1>
            </div>

            <!-- Pill Badges (matching screenshot chips) -->
            <div class="flex items-center gap-3 pt-2">
              <div class="inline-flex items-center gap-2 px-3 py-1.5 rounded-full bg-slate-900/60 backdrop-blur-md border border-white/20 text-xs text-white shadow-sm">
                <span class="w-2.5 h-2.5 rounded-full bg-gradient-to-r from-amber-400 to-orange-500"></span>
                <span class="font-bold text-[11px]">5 Stage Pipeline</span>
              </div>
              <div class="inline-flex items-center gap-2 px-3 py-1.5 rounded-full bg-slate-900/60 backdrop-blur-md border border-white/20 text-xs text-white shadow-sm">
                <span class="w-2.5 h-2.5 rounded-full bg-gradient-to-r from-cyan-400 to-blue-500"></span>
                <span class="font-bold text-[11px]">Technical Velocity</span>
              </div>
            </div>
          </div>

          <!-- Right Floating Frosted Glass Sign-In Card -->
          <div class="w-full lg:w-[420px] p-8 md:p-10 rounded-3xl kaleido-glass text-white space-y-6">
            <div class="space-y-1">
              <h2 class="text-2xl font-black text-white">Sign in</h2>
              <p class="text-xs text-slate-300/80">Good to see you again.</p>
            </div>

            <form onsubmit="handleLogin(event)" class="space-y-4 pt-1">
              <!-- Username / Email Input -->
              <div class="space-y-1.5">
                <label class="text-[11px] font-semibold text-slate-200">Username</label>
                <input id="loginUsername" type="text" required placeholder="admin" class="w-full text-xs px-4 py-3 rounded-xl bg-white/10 border border-white/20 text-white placeholder-slate-400 outline-none focus:ring-2 focus:ring-white/50 focus:bg-white/15 transition shadow-inner">
              </div>

              <!-- Password Input -->
              <div class="space-y-1.5">
                <label class="text-[11px] font-semibold text-slate-200">Password</label>
                <div class="relative">
                  <input id="loginPassword" type="password" required placeholder="••••••••" class="w-full text-xs px-4 py-3 rounded-xl bg-white/10 border border-white/20 text-white placeholder-slate-400 outline-none focus:ring-2 focus:ring-white/50 focus:bg-white/15 transition shadow-inner pr-10">
                  <button type="button" onclick="togglePasswordVisibility()" class="absolute right-3.5 top-3.5 text-slate-400 hover:text-white text-xs">
                    <i id="passwordEyeIcon" class="fa-regular fa-eye"></i>
                  </button>
                </div>
              </div>

              <!-- Remember Me & Forgot Password -->
              <div class="flex items-center justify-between text-xs pt-1">
                <label class="flex items-center gap-2 cursor-pointer text-slate-300">
                  <input type="checkbox" checked class="w-3.5 h-3.5 rounded bg-white/20 border-white/30 text-orange-500 focus:ring-0">
                  <span class="text-[11px]">Remember me</span>
                </label>
                <a href="#" onclick="alert('Default login credentials are: admin / admin123')" class="text-[11px] text-slate-300 hover:text-white underline">Forgot password?</a>
              </div>

              <!-- Error Box -->
              <div id="loginError" class="hidden text-xs text-rose-200 bg-rose-500/25 border border-rose-500/40 p-2.5 rounded-xl text-center font-medium"></div>

              <!-- Submit Button (White Pill Button from screenshot) -->
              <button type="submit" class="w-full py-3 bg-white hover:bg-slate-100 text-slate-900 rounded-xl text-xs font-bold shadow-xl hover:-translate-y-0.5 transition flex items-center justify-center gap-2">
                <span>Sign in</span>
                <i class="fa-solid fa-arrow-right text-xs"></i>
              </button>
            </form>

            <div class="pt-2 text-center text-slate-400 text-[10px]">
              Default login: <strong class="text-white">admin</strong> / <strong class="text-white">admin123</strong>
            </div>
          </div>

        </div>

      </div>

      <!-- ================= MAIN APP SHELL ================= -->
      <div id="appShell" class="hidden max-w-7xl mx-auto p-4 md:p-6 space-y-6">

        <!-- Top Header Navigation Bar (without phrase, simplified) -->
        <header class="glass-card p-4 md:p-5 rounded-2xl shadow-sm flex flex-wrap justify-between items-center gap-4 border border-slate-200/90">
          <div class="flex items-center gap-3.5">
            <div class="w-12 h-12 rounded-2xl bg-gradient-to-tr from-orange-400 to-indigo-600 text-white flex items-center justify-center font-bold text-xl shadow-md shadow-orange-400/20 ring-4 ring-orange-400/10">
              <i class="fa-solid fa-compass-drafting"></i>
            </div>
            <div>
              <div class="flex items-center gap-2">
                <h1 class="text-xl font-black bg-gradient-to-r from-slate-900 via-indigo-950 to-orange-500 bg-clip-text text-transparent">
                  Presales Activities Hub
                </h1>
                <span id="roleBadge" class="text-[10px] font-extrabold px-2 py-0.5 rounded-full uppercase bg-orange-50 text-orange-600 border border-orange-200">ADMIN</span>
              </div>
            </div>
          </div>

          <!-- Header Actions: Navigation Tabs & Admin Tools -->
          <div class="flex items-center flex-wrap gap-2.5">
            <button onclick="switchTab('exec')" id="tabBtnExec" class="px-4 py-2 rounded-xl text-xs font-bold bg-gradient-to-r from-orange-400 to-amber-500 text-white shadow-md shadow-orange-400/20 hover:-translate-y-0.5 transition">
              <i class="fa-solid fa-chart-pie mr-1.5"></i> Dashboard
            </button>
            <button onclick="switchTab('stage')" id="tabBtnStage" class="px-4 py-2 rounded-xl text-xs font-bold bg-slate-100 hover:bg-orange-50 hover:text-orange-600 text-slate-700 border border-slate-200 hover:-translate-y-0.5 transition">
              <i class="fa-solid fa-cubes-stacked mr-1.5 text-indigo-600"></i> Workstream
            </button>
            <button onclick="switchTab('audit')" id="tabBtnAudit" class="px-4 py-2 rounded-xl text-xs font-bold bg-slate-100 hover:bg-blue-50 hover:text-blue-600 text-slate-700 border border-slate-200 hover:-translate-y-0.5 transition">
              <i class="fa-solid fa-clock-rotate-left mr-1.5 text-slate-400"></i> Audit History
            </button>

            <!-- Admin User Management Button -->
            <button onclick="openUsersModal()" id="btnUsersMgmt" class="px-3 py-2 rounded-xl text-xs font-bold bg-slate-100 hover:bg-slate-200 text-slate-700 border border-slate-200 transition" title="User Management">
              <i class="fa-solid fa-users-gear"></i>
            </button>

            <!-- Sign Out -->
            <button onclick="handleLogout()" class="px-3 py-2 rounded-xl text-xs font-bold bg-rose-50 hover:bg-rose-100 text-rose-600 border border-rose-200 transition" title="Sign Out">
              <i class="fa-solid fa-power-off"></i>
            </button>
          </div>
        </header>

        <!-- ================= TAB 1: EXECUTIVE DASHBOARD ================= -->
        <div id="viewExec" class="space-y-6">
          
          <div class="flex flex-wrap justify-between items-center glass-card p-5 rounded-2xl shadow-sm gap-4 border-l-4 border-l-orange-400">
            <div>
              <div class="flex items-center gap-2">
                <span class="w-2.5 h-2.5 rounded-full bg-emerald-500 animate-pulse"></span>
                <h2 class="text-base font-bold text-slate-900">Executive Overview</h2>
              </div>
              <p class="text-xs text-slate-500 mt-0.5">Click any metric card or stage tile to open its focused detail window.</p>
            </div>
            <div class="flex items-center gap-2.5">
              <button onclick="window.location.href='/api/reports/export/pdf'" class="flex items-center gap-2 px-3.5 py-2 bg-gradient-to-r from-rose-50 to-orange-50 hover:from-rose-100 hover:to-orange-100 text-rose-700 border border-rose-200 text-xs font-bold rounded-xl shadow-sm hover:-translate-y-0.5 transition">
                <i class="fa-solid fa-file-pdf text-rose-600"></i> Export to PDF
              </button>
              <button onclick="window.location.href='/api/reports/export/pptx'" class="flex items-center gap-2 px-3.5 py-2 bg-gradient-to-r from-amber-50 to-orange-50 hover:from-amber-100 hover:to-orange-100 text-orange-800 border border-orange-200 text-xs font-bold rounded-xl shadow-sm hover:-translate-y-0.5 transition">
                <i class="fa-solid fa-file-powerpoint text-orange-600"></i> Export to PPTX
              </button>
            </div>
          </div>

          <!-- Highlighted High-Visibility KPI Cards -->
          <div class="grid grid-cols-2 md:grid-cols-4 gap-4">
            <!-- 1. Engaged Accounts -->
            <div onclick="openDrilldownModal('engaged')" class="glass-card tile-card p-5 rounded-2xl border-2 border-indigo-200 hover:border-indigo-400 transition cursor-pointer group">
              <div class="flex justify-between items-start">
                <span class="text-[11px] font-extrabold uppercase tracking-wider text-indigo-700">Engaged Accounts</span>
                <div class="w-9 h-9 rounded-xl bg-indigo-50 text-indigo-600 flex items-center justify-center text-xs group-hover:bg-indigo-600 group-hover:text-white transition shadow-sm">
                  <i class="fa-solid fa-building-user text-sm"></i>
                </div>
              </div>
              <div id="kpiCustCount" class="text-3xl font-black text-slate-900 mt-2">0</div>
              <span class="text-[10px] text-indigo-600 font-bold mt-2 flex items-center gap-1 group-hover:translate-x-1 transition">
                Inspect accounts &rarr;
              </span>
            </div>

            <!-- 2. Active Deals -->
            <div onclick="openDrilldownModal('active_opps')" class="glass-card tile-card p-5 rounded-2xl border-2 border-orange-200 hover:border-orange-400 transition cursor-pointer group">
              <div class="flex justify-between items-start">
                <span class="text-[11px] font-extrabold uppercase tracking-wider text-orange-600">Active Deals</span>
                <div class="w-9 h-9 rounded-xl bg-orange-50 text-orange-500 flex items-center justify-center text-xs group-hover:bg-orange-500 group-hover:text-white transition shadow-sm">
                  <i class="fa-solid fa-fire text-sm"></i>
                </div>
              </div>
              <div id="kpiOppCount" class="text-3xl font-black text-orange-600 mt-2">0</div>
              <span class="text-[10px] text-orange-600 font-bold mt-2 flex items-center gap-1 group-hover:translate-x-1 transition">
                View pipeline &rarr;
              </span>
            </div>

            <!-- 3. Urgent Focus -->
            <div onclick="openDrilldownModal('urgent')" class="glass-card tile-card p-5 rounded-2xl border-2 border-rose-200 hover:border-rose-400 transition cursor-pointer group">
              <div class="flex justify-between items-start">
                <span class="text-[11px] font-extrabold uppercase tracking-wider text-rose-600">Urgent Focus</span>
                <div class="w-9 h-9 rounded-xl bg-rose-50 text-rose-600 flex items-center justify-center text-xs group-hover:bg-rose-600 group-hover:text-white transition shadow-sm">
                  <i class="fa-solid fa-triangle-exclamation text-sm"></i>
                </div>
              </div>
              <div id="kpiUrgentCount" class="text-3xl font-black text-rose-600 mt-2">0</div>
              <span class="text-[10px] text-rose-600 font-bold mt-2 flex items-center gap-1 group-hover:translate-x-1 transition">
                Action required &rarr;
              </span>
            </div>

            <!-- 4. Completed Deliverables -->
            <div onclick="openDrilldownModal('completed')" class="glass-card tile-card p-5 rounded-2xl border-2 border-emerald-200 hover:border-emerald-400 transition cursor-pointer group">
              <div class="flex justify-between items-start">
                <span class="text-[11px] font-extrabold uppercase tracking-wider text-emerald-600">Completed Tasks</span>
                <div class="w-9 h-9 rounded-xl bg-emerald-50 text-emerald-600 flex items-center justify-center text-xs group-hover:bg-emerald-600 group-hover:text-white transition shadow-sm">
                  <i class="fa-solid fa-circle-check text-sm"></i>
                </div>
              </div>
              <div id="kpiCompletedTasks" class="text-3xl font-black text-emerald-600 mt-2">0</div>
              <span class="text-[10px] text-emerald-600 font-bold mt-2 flex items-center gap-1 group-hover:translate-x-1 transition">
                Audit delivery &rarr;
              </span>
            </div>
          </div>

          <!-- Status & Stage Breakdown Cards -->
          <div class="grid grid-cols-1 md:grid-cols-2 gap-6">
            <div class="glass-card tile-card p-5 rounded-2xl border-2 border-slate-200/90 space-y-3">
              <h3 class="text-xs font-extrabold uppercase tracking-wider text-slate-700 flex items-center gap-2">
                <i class="fa-solid fa-bars-progress text-indigo-600"></i> Tasks by Status (Click to inspect)
              </h3>
              <div class="grid grid-cols-2 sm:grid-cols-4 gap-2.5" id="statusTilesContainer"></div>
            </div>

            <div class="glass-card tile-card p-5 rounded-2xl border-2 border-slate-200/90 space-y-3">
              <h3 class="text-xs font-extrabold uppercase tracking-wider text-slate-700 flex items-center gap-2">
                <i class="fa-solid fa-layer-group text-orange-500"></i> Tasks by Stage (Click to inspect)
              </h3>
              <div class="grid grid-cols-2 sm:grid-cols-4 gap-2.5" id="stageTilesContainer"></div>
            </div>
          </div>

          <!-- Opportunities Portfolio Matrix -->
          <div class="glass-card tile-card rounded-2xl overflow-hidden border-2 border-slate-200/90">
            <div class="p-4 md:p-5 border-b border-slate-100 flex justify-between items-center flex-wrap gap-2">
              <div class="flex items-center gap-2.5">
                <div class="w-8 h-8 rounded-lg bg-orange-100 text-orange-600 flex items-center justify-center font-bold text-xs">
                  <i class="fa-solid fa-briefcase"></i>
                </div>
                <div>
                  <h3 class="text-sm font-bold text-slate-900">Active Opportunities Portfolio</h3>
                  <p class="text-[11px] text-slate-400">Click column headers to sort</p>
                </div>
              </div>
              <div id="sortStatusBadge" class="text-xs bg-orange-50 text-orange-700 px-3 py-1 rounded-full font-bold border border-orange-200">
                Sorted by Priority: High &rarr; Low
              </div>
            </div>
            <div class="overflow-x-auto">
              <table class="w-full text-left text-xs">
                <thead class="bg-slate-50 text-slate-500 text-[10px] uppercase tracking-wider border-b border-slate-200 select-none">
                  <tr>
                    <th onclick="toggleSort('customer')" class="py-3 px-4 cursor-pointer hover:bg-slate-100 transition">
                      Account <i id="sortIconCustomer" class="fa-solid fa-sort ml-1"></i>
                    </th>
                    <th class="py-3 px-4">Opportunity</th>
                    <th class="py-3 px-4">Target End Date</th>
                    <th onclick="toggleSort('priority')" class="py-3 px-4 cursor-pointer hover:bg-slate-100 transition">
                      Priority <i id="sortIconPriority" class="fa-solid fa-sort-down ml-1 text-orange-500"></i>
                    </th>
                    <th onclick="toggleSort('status')" class="py-3 px-4 cursor-pointer hover:bg-slate-100 transition">
                      Status <i id="sortIconStatus" class="fa-solid fa-sort ml-1"></i>
                    </th>
                    <th class="py-3 px-4 text-right">Actions</th>
                  </tr>
                </thead>
                <tbody id="execMatrixBody" class="divide-y divide-slate-100 font-medium"></tbody>
              </table>
            </div>
          </div>

          <!-- Scheduled Tasks Calendar -->
          <div class="glass-card tile-card rounded-2xl p-5 md:p-6 space-y-4 border-2 border-slate-200/90">
            <div class="flex flex-wrap justify-between items-center gap-2 pb-3 border-b border-slate-100">
              <div class="flex items-center gap-2">
                <div class="w-8 h-8 rounded-lg bg-indigo-100 text-indigo-600 flex items-center justify-center font-bold text-xs">
                  <i class="fa-regular fa-calendar-days"></i>
                </div>
                <div>
                  <h3 class="text-sm font-bold text-slate-900">Scheduled Engagement Calendar</h3>
                  <p class="text-[11px] text-slate-400">View tasks scheduled across all client engagements</p>
                </div>
              </div>
              <div class="flex items-center gap-1.5 bg-slate-100 p-1 rounded-xl border border-slate-200">
                <button onclick="changeCalMonth(-1)" class="p-1.5 hover:bg-white rounded-lg text-slate-600 transition"><i class="fa-solid fa-chevron-left text-xs"></i></button>
                <span id="calMonthYearLabel" class="text-xs font-bold text-slate-800 w-32 text-center">Month Year</span>
                <button onclick="changeCalMonth(1)" class="p-1.5 hover:bg-white rounded-lg text-slate-600 transition"><i class="fa-solid fa-chevron-right text-xs"></i></button>
              </div>
            </div>

            <div class="grid grid-cols-1 lg:grid-cols-3 gap-6">
              <div class="lg:col-span-2">
                <div class="grid grid-cols-7 gap-1 text-center text-[10px] font-bold text-slate-400 uppercase py-2">
                  <div>Sun</div><div>Mon</div><div>Tue</div><div>Wed</div><div>Thu</div><div>Fri</div><div>Sat</div>
                </div>
                <div id="calendarDaysGrid" class="grid grid-cols-7 gap-1.5"></div>
              </div>

              <div class="bg-slate-50/80 border border-slate-200 rounded-2xl p-4 flex flex-col h-full min-h-[320px]">
                <div class="pb-2.5 border-b border-slate-200 flex justify-between items-center">
                  <span class="text-xs font-bold text-slate-900 flex items-center gap-1.5" id="agendaSelectedDateLabel">
                    <i class="fa-solid fa-calendar-day text-orange-500"></i> Scheduled Tasks
                  </span>
                  <span class="text-[10px] bg-orange-100 text-orange-700 px-2.5 py-0.5 rounded-full font-bold border border-orange-200" id="agendaTaskCountBadge">0 Tasks</span>
                </div>
                <div id="agendaTaskList" class="flex-1 overflow-y-auto space-y-2 mt-3 pr-1 text-xs custom-scrollbar"></div>
              </div>
            </div>
          </div>

        </div>

        <!-- ================= TAB 2: MASTER-DETAIL STAGE WORKSTREAM ================= -->
        <div id="viewStage" class="hidden">
          <div class="grid grid-cols-1 lg:grid-cols-12 gap-5 items-start">
            
            <!-- LEFT PANEL: Account & Deal Navigator -->
            <div class="lg:col-span-4 glass-card tile-card rounded-2xl p-4 space-y-4 border-2 border-slate-200/90">
              <div class="flex justify-between items-center pb-3 border-b border-slate-100">
                <div>
                  <h3 class="text-sm font-bold text-slate-900 flex items-center gap-1.5">
                    <i class="fa-solid fa-sitemap text-indigo-600"></i> Accounts & Deals
                  </h3>
                  <p class="text-[11px] text-slate-400">Select any opportunity to manage</p>
                </div>
                <button onclick="openWizardModal()" id="btnNewEngagement" class="px-3 py-1.5 bg-gradient-to-r from-orange-400 to-amber-500 hover:from-orange-500 hover:to-amber-600 text-white rounded-xl text-xs font-bold shadow-md shadow-orange-400/20 transition flex items-center gap-1.5">
                  <i class="fa-solid fa-plus text-[10px]"></i> New Deal
                </button>
              </div>

              <!-- Search -->
              <div class="relative">
                <i class="fa-solid fa-magnifying-glass absolute left-3 top-2.5 text-slate-400 text-xs"></i>
                <input id="navSearchInput" oninput="renderNavigatorTree()" type="text" placeholder="Filter customer or deal..." class="w-full text-xs pl-8 pr-3 py-2 bg-slate-50 border border-slate-200 rounded-xl text-slate-800 outline-none focus:ring-2 focus:ring-orange-400 transition">
              </div>

              <div id="navigatorTreeList" class="space-y-3 max-h-[640px] overflow-y-auto custom-scrollbar pr-1"></div>
            </div>

            <!-- RIGHT PANEL: Focus Workspace & 5 Stages -->
            <div class="lg:col-span-8 space-y-4">
              
              <!-- Sticky Active Context Banner -->
              <div id="oppActiveBanner" class="glass-card tile-card p-5 rounded-2xl flex flex-wrap justify-between items-center gap-3 relative overflow-hidden border-2 border-slate-200/90">
                <div class="absolute left-0 top-0 bottom-0 w-1.5 bg-gradient-to-b from-orange-400 to-indigo-600"></div>
                <div class="pl-2">
                  <div class="flex items-center gap-2">
                    <span id="bannerCustName" class="text-xs uppercase font-black text-orange-500 tracking-wider">Account</span>
                    <button onclick="openEditCustomerModal()" id="btnEditCust" class="text-[10px] text-slate-400 hover:text-indigo-600" title="Edit Customer"><i class="fa-solid fa-pen"></i></button>
                  </div>
                  <h2 id="bannerOppTitle" class="text-lg font-black text-slate-900 mt-0.5">Select a deal to start</h2>
                  <div class="flex items-center gap-3 text-xs text-slate-500 mt-1">
                    <span>Target End: <strong id="bannerOppDate" class="text-slate-800">-</strong></span>
                    <span>&bull;</span>
                    <span>Priority: <strong id="bannerOppPriority" class="text-slate-800">-</strong></span>
                    <span>&bull;</span>
                    <span>Status: <strong id="bannerOppStatus" class="text-slate-800">-</strong></span>
                  </div>
                </div>
                <div class="flex items-center gap-2">
                  <button onclick="openEditCurrentOppModal()" id="btnEditOpp" class="px-3.5 py-2 bg-slate-100 hover:bg-orange-50 hover:text-orange-700 text-slate-700 rounded-xl text-xs font-bold border border-slate-200 transition">
                    <i class="fa-solid fa-pen-to-square mr-1.5 text-orange-400"></i> Edit Deal
                  </button>
                </div>
              </div>

              <!-- 5 Stage Gate Navigator Pills -->
              <div class="glass-card tile-card p-2.5 rounded-2xl flex flex-wrap gap-1.5 border-2 border-slate-200/90" id="stagePillsContainer"></div>

              <!-- Active Stage Workspace -->
              <div id="stageContentArea" class="glass-card tile-card p-6 rounded-2xl min-h-[440px] border-2 border-slate-200/90"></div>

            </div>

          </div>
        </div>

        <!-- ================= TAB 3: AUDIT HISTORY ================= -->
        <div id="viewAudit" class="hidden space-y-4">
          <div class="glass-card tile-card p-5 rounded-2xl flex justify-between items-center border-2 border-slate-200/90">
            <div>
              <h2 class="text-base font-bold text-slate-900 flex items-center gap-2">
                <i class="fa-solid fa-shield-halved text-orange-500"></i> Change Audit Trail
              </h2>
              <p class="text-xs text-slate-500 mt-0.5">Immutable record of modifications made across client accounts.</p>
            </div>
          </div>
          <div class="glass-card tile-card rounded-2xl overflow-hidden border-2 border-slate-200/90">
            <div class="overflow-x-auto">
              <table class="w-full text-left text-xs">
                <thead class="bg-slate-50 text-slate-500 text-[10px] uppercase border-b border-slate-200">
                  <tr>
                    <th class="py-3 px-4">Timestamp</th>
                    <th class="py-3 px-4">Event Type</th>
                    <th class="py-3 px-4">Customer</th>
                    <th class="py-3 px-4">Opportunity</th>
                    <th class="py-3 px-4">Details</th>
                  </tr>
                </thead>
                <tbody id="auditTableBody" class="divide-y divide-slate-100 font-medium"></tbody>
              </table>
            </div>
          </div>
        </div>

      </div>

      <!-- ================= MODALS ================= -->

      <!-- Users Management Modal -->
      <div id="modalUsers" class="hidden fixed inset-0 bg-slate-900/40 backdrop-blur-sm z-50 flex items-center justify-center p-4">
        <div class="bg-white w-full max-w-xl p-6 rounded-2xl shadow-2xl border border-slate-200 space-y-4">
          <div class="flex justify-between items-center pb-2 border-b border-slate-100">
            <div>
              <h3 class="text-base font-bold text-slate-900">User & Role Management</h3>
              <p class="text-xs text-slate-500">Manage user logins and assign access permissions</p>
            </div>
            <button onclick="closeModal('modalUsers')" class="text-slate-400 hover:text-slate-600 text-sm"><i class="fa-solid fa-xmark"></i></button>
          </div>

          <form onsubmit="handleCreateUser(event)" class="bg-slate-50 p-4 rounded-xl border border-slate-200 space-y-3">
            <span class="text-xs font-bold text-orange-600 uppercase tracking-wider block">Create New User</span>
            <div class="grid grid-cols-2 gap-2">
              <input id="newUsername" type="text" required placeholder="Username" class="text-xs p-2 bg-white border border-slate-200 rounded-lg text-slate-800 outline-none">
              <input id="newPassword" type="password" required placeholder="Password" class="text-xs p-2 bg-white border border-slate-200 rounded-lg text-slate-800 outline-none">
              <input id="newFullName" type="text" required placeholder="Full Name" class="text-xs p-2 bg-white border border-slate-200 rounded-lg text-slate-800 outline-none">
              <select id="newRole" class="text-xs p-2 bg-white border border-slate-200 rounded-lg text-slate-800 outline-none">
                <option value="view_only">View Only</option>
                <option value="admin">Administrator (Full Access)</option>
              </select>
            </div>
            <div class="flex justify-end">
              <button type="submit" class="px-3.5 py-1.5 bg-orange-500 hover:bg-orange-600 text-white font-bold text-xs rounded-lg transition">
                Add User
              </button>
            </div>
          </form>

          <div class="space-y-2">
            <span class="text-xs font-bold text-slate-400 uppercase tracking-wider block">Active Users</span>
            <div id="usersListTable" class="max-h-48 overflow-y-auto space-y-1.5 pr-1 custom-scrollbar text-xs"></div>
          </div>
        </div>
      </div>

      <!-- Unified Engagement Wizard Modal -->
      <div id="modalWizard" class="hidden fixed inset-0 bg-slate-900/40 backdrop-blur-sm z-50 flex items-center justify-center p-4">
        <div class="bg-white w-full max-w-lg p-6 rounded-2xl shadow-2xl border border-slate-200 space-y-4">
          <div class="flex justify-between items-center pb-2 border-b border-slate-100">
            <div class="flex items-center gap-2">
              <div class="w-8 h-8 rounded-xl bg-orange-100 text-orange-600 flex items-center justify-center text-xs font-bold">
                <i class="fa-solid fa-wand-magic-sparkles"></i>
              </div>
              <div>
                <h3 class="text-base font-bold text-slate-900">New Engagement Wizard</h3>
                <p class="text-xs text-slate-400">Launch a new presales workstream</p>
              </div>
            </div>
            <button onclick="closeModal('modalWizard')" class="text-slate-400 hover:text-slate-600 text-sm"><i class="fa-solid fa-xmark"></i></button>
          </div>

          <div class="space-y-1.5">
            <label class="text-[11px] font-bold uppercase tracking-wider text-slate-500">Step 1: Choose Account Mode</label>
            <div class="grid grid-cols-2 gap-2 p-1.5 bg-slate-100 rounded-xl">
              <button type="button" onclick="setWizardMode('existing')" id="wizBtnExisting" class="py-2 text-xs font-bold rounded-lg bg-white text-orange-600 shadow-sm transition">
                Existing Account
              </button>
              <button type="button" onclick="setWizardMode('new')" id="wizBtnNew" class="py-2 text-xs font-bold rounded-lg text-slate-600 hover:text-slate-900 transition">
                + New Customer
              </button>
            </div>
          </div>

          <div id="wizExistingGroup" class="space-y-1">
            <label class="text-[11px] font-semibold text-slate-600">Select Customer Account</label>
            <select id="wizCustSelect" class="w-full text-xs border border-slate-200 bg-white rounded-xl p-2.5 text-slate-800 outline-none focus:ring-2 focus:ring-orange-400"></select>
          </div>

          <div id="wizNewGroup" class="hidden space-y-2.5">
            <div>
              <label class="text-[11px] font-semibold text-slate-600">Customer Name</label>
              <input id="wizNewCustName" type="text" placeholder="e.g., Apex Healthcare Systems" class="w-full text-xs border border-slate-200 rounded-xl p-2.5 text-slate-800 outline-none focus:ring-2 focus:ring-orange-400">
            </div>
            <div>
              <label class="text-[11px] font-semibold text-slate-600">Industry Sector</label>
              <input id="wizNewCustIndustry" type="text" placeholder="e.g., Banking, Healthcare" class="w-full text-xs border border-slate-200 rounded-xl p-2.5 text-slate-800 outline-none focus:ring-2 focus:ring-orange-400">
            </div>
          </div>

          <div class="pt-2 border-t border-slate-100 space-y-3">
            <label class="text-[11px] font-bold uppercase tracking-wider text-slate-500">Step 2: Opportunity Scope</label>
            <div>
              <label class="text-[11px] font-semibold text-slate-600">Opportunity Title</label>
              <input id="wizOppTitle" type="text" placeholder="e.g., Azure Identity & Defender Rollout" class="w-full text-xs border border-slate-200 rounded-xl p-2.5 text-slate-800 outline-none focus:ring-2 focus:ring-orange-400">
            </div>

            <div class="grid grid-cols-3 gap-2.5">
              <div>
                <label class="text-[11px] font-semibold text-slate-600">Priority</label>
                <select id="wizOppPriority" class="w-full text-xs border border-slate-200 rounded-xl p-2 text-slate-800 outline-none">
                  <option>Medium</option>
                  <option>High</option>
                  <option>Urgent</option>
                  <option>Low</option>
                </select>
              </div>
              <div>
                <label class="text-[11px] font-semibold text-slate-600">Status</label>
                <select id="wizOppStatus" class="w-full text-xs border border-slate-200 rounded-xl p-2 text-slate-800 outline-none">
                  <option>New</option>
                  <option>On progress</option>
                  <option>Hold</option>
                  <option>Completed</option>
                </select>
              </div>
              <div>
                <label class="text-[11px] font-semibold text-slate-600">Target End Date</label>
                <input id="wizOppEnd" type="date" class="w-full text-xs border border-slate-200 rounded-xl p-2 text-slate-800 outline-none">
              </div>
            </div>
          </div>

          <div class="flex justify-end gap-2.5 pt-3 border-t border-slate-100">
            <button onclick="closeModal('modalWizard')" class="px-4 py-2 text-xs text-slate-600 rounded-xl hover:bg-slate-100">Cancel</button>
            <button onclick="submitWizard()" class="px-5 py-2 text-xs font-bold bg-orange-500 hover:bg-orange-600 text-white rounded-xl shadow transition">
              Create Engagement
            </button>
          </div>
        </div>
      </div>

      <!-- Drilldown Details Modal for Tiles -->
      <div id="modalTileDetails" class="hidden fixed inset-0 bg-slate-900/40 backdrop-blur-sm z-50 flex items-center justify-center p-4">
        <div class="bg-white w-full max-w-3xl p-6 rounded-2xl shadow-2xl border border-slate-200 space-y-4 max-h-[85vh] flex flex-col">
          <div class="flex justify-between items-center pb-3 border-b border-slate-100">
            <div>
              <h3 id="tileModalTitle" class="text-base font-bold text-slate-900">Details</h3>
              <p id="tileModalSubtitle" class="text-xs text-slate-500">Drill-down inspection</p>
            </div>
            <button onclick="closeModal('modalTileDetails')" class="text-slate-400 hover:text-slate-600 text-sm"><i class="fa-solid fa-xmark"></i></button>
          </div>
          <div id="tileModalContent" class="flex-1 overflow-y-auto space-y-2.5 pr-1 text-xs custom-scrollbar"></div>
        </div>
      </div>

      <!-- Edit Customer Modal -->
      <div id="modalEditCustomer" class="hidden fixed inset-0 bg-slate-900/40 backdrop-blur-sm z-50 flex items-center justify-center p-4">
        <div class="bg-white w-full max-w-sm p-6 rounded-2xl shadow-2xl border border-slate-200 space-y-4">
          <h3 class="text-base font-bold text-slate-900">Edit Customer Account</h3>
          <input type="hidden" id="editCustId">
          <div>
            <label class="text-[11px] font-semibold text-slate-600 block mb-1">Customer / Account Name</label>
            <input id="editCustName" type="text" class="w-full text-xs border border-slate-200 rounded-xl p-2.5 text-slate-800 outline-none">
          </div>
          <div>
            <label class="text-[11px] font-semibold text-slate-600 block mb-1">Industry Sector</label>
            <input id="editCustIndustry" type="text" class="w-full text-xs border border-slate-200 rounded-xl p-2.5 text-slate-800 outline-none">
          </div>
          <div class="flex justify-end gap-2 pt-2">
            <button onclick="closeModal('modalEditCustomer')" class="px-3.5 py-1.5 text-xs text-slate-600">Cancel</button>
            <button onclick="submitEditCustomer()" class="px-4 py-1.5 text-xs font-bold bg-orange-500 text-white rounded-xl">Update</button>
          </div>
        </div>
      </div>

      <!-- Edit Opportunity Modal -->
      <div id="modalEditOpp" class="hidden fixed inset-0 bg-slate-900/40 backdrop-blur-sm z-50 flex items-center justify-center p-4">
        <div class="bg-white w-full max-w-md p-6 rounded-2xl shadow-2xl border border-slate-200 space-y-4">
          <h3 class="text-base font-bold text-slate-900">Edit Opportunity Details</h3>
          <input type="hidden" id="editOppId">
          <div>
            <label class="text-[11px] font-semibold text-slate-600 block mb-1">Opportunity Title</label>
            <input id="editOppTitle" type="text" class="w-full text-xs border border-slate-200 rounded-xl p-2.5 text-slate-800 outline-none">
          </div>
          <div>
            <label class="text-[11px] font-semibold text-slate-600 block mb-1">Target End Date</label>
            <input id="editOppEnd" type="date" class="w-full text-xs border border-slate-200 rounded-xl p-2 text-slate-800 outline-none">
          </div>
          <div class="grid grid-cols-2 gap-3">
            <div>
              <label class="text-[11px] font-semibold text-slate-600 block mb-1">Priority</label>
              <select id="editOppPriority" class="w-full text-xs border border-slate-200 rounded-xl p-2 text-slate-800 outline-none">
                <option>Low</option>
                <option>Medium</option>
                <option>High</option>
                <option>Urgent</option>
              </select>
            </div>
            <div>
              <label class="text-[11px] font-semibold text-slate-600 block mb-1">Status</label>
              <select id="editOppStatus" class="w-full text-xs border border-slate-200 rounded-xl p-2 text-slate-800 outline-none">
                <option>New</option>
                <option>On progress</option>
                <option>Hold</option>
                <option>Completed</option>
              </select>
            </div>
          </div>
          <div class="flex justify-end gap-2 pt-2">
            <button onclick="closeModal('modalEditOpp')" class="px-3.5 py-1.5 text-xs text-slate-600">Cancel</button>
            <button onclick="submitEditOpportunity()" class="px-4 py-1.5 text-xs font-bold bg-orange-500 text-white rounded-xl">Save</button>
          </div>
        </div>
      </div>

      <!-- Add / Edit Task Modal -->
      <div id="modalTask" class="hidden fixed inset-0 bg-slate-900/40 backdrop-blur-sm z-50 flex items-center justify-center p-4">
        <div class="bg-white w-full max-w-md p-6 rounded-2xl shadow-2xl border border-slate-200 space-y-4">
          <h3 id="modalTaskTitle" class="text-base font-bold text-slate-900">Task Details</h3>
          <input type="hidden" id="taskEditId">
          <input id="inputTaskTitle" type="text" placeholder="Task description..." class="w-full text-xs border border-slate-200 rounded-xl p-2.5 text-slate-800 outline-none">
          <div class="grid grid-cols-2 gap-3">
            <div>
              <label class="text-[11px] font-semibold text-slate-600 block mb-1">Start Date</label>
              <input id="inputTaskStart" type="date" class="w-full text-xs border border-slate-200 rounded-xl p-2 text-slate-800 outline-none">
            </div>
            <div>
              <label class="text-[11px] font-semibold text-slate-600 block mb-1">End Date</label>
              <input id="inputTaskEnd" type="date" class="w-full text-xs border border-slate-200 rounded-xl p-2 text-slate-800 outline-none">
            </div>
          </div>
          <div>
            <label class="text-[11px] font-semibold text-slate-600 block mb-1">Status</label>
            <select id="inputTaskStatus" class="w-full text-xs border border-slate-200 rounded-xl p-2.5 text-slate-800 outline-none">
              <option>New</option>
              <option>On going</option>
              <option>Hold</option>
              <option>Completed</option>
            </select>
          </div>
          <div class="flex justify-end gap-2 pt-2">
            <button onclick="closeModal('modalTask')" class="px-3.5 py-1.5 text-xs text-slate-600">Cancel</button>
            <button onclick="submitTaskModal()" class="px-4 py-1.5 text-xs font-bold bg-orange-500 text-white rounded-xl">Save Task</button>
          </div>
        </div>
      </div>

      <!-- Daily Comment Log Modal -->
      <div id="modalComment" class="hidden fixed inset-0 bg-slate-900/40 backdrop-blur-sm z-50 flex items-center justify-center p-4">
        <div class="bg-white w-full max-w-lg p-6 rounded-2xl shadow-2xl border border-slate-200 flex flex-col max-h-[85vh]">
          <div class="flex justify-between items-start pb-3 border-b border-slate-100">
            <div>
              <span class="text-[10px] font-bold uppercase tracking-wider text-orange-500">Daily Updates Log</span>
              <h3 id="commentModalTaskTitle" class="text-xs font-bold text-slate-900 mt-0.5 line-clamp-1">-</h3>
            </div>
            <button onclick="closeModal('modalComment')" class="text-slate-400 hover:text-slate-600 text-xs"><i class="fa-solid fa-xmark"></i></button>
          </div>
          <div id="commentList" class="flex-1 overflow-y-auto my-4 space-y-2 pr-1 text-xs custom-scrollbar"></div>
          <div id="commentFormContainer" class="pt-3 border-t border-slate-100 space-y-2">
            <textarea id="commentInputText" rows="2" placeholder="Post today's progress, blocker, or update..." class="w-full text-xs border border-slate-200 rounded-xl p-2.5 text-slate-800 outline-none focus:ring-2 focus:ring-orange-400"></textarea>
            <div class="flex justify-end">
              <button onclick="submitTaskComment()" class="px-4 py-2 text-xs font-bold bg-orange-500 hover:bg-orange-600 text-white rounded-xl shadow transition">
                <i class="fa-solid fa-paper-plane mr-1 text-[10px]"></i> Post Update
              </button>
            </div>
          </div>
        </div>
      </div>

      <script>
        const STAGES = [
          { id: 'Stage 01 Strategy', name: 'Stage 01 Strategy', icon: 'fa-lightbulb', isStrategy: true },
          { id: 'Stage 02 Assessment', name: 'Stage 02 Assessment', icon: 'fa-magnifying-glass-chart', isStrategy: false },
          { id: 'Stage 03 POC', name: 'Stage 03 POC', icon: 'fa-flask', isStrategy: false },
          { id: 'Stage 04 Designing', name: 'Stage 04 Designing', icon: 'fa-network-wired', isStrategy: false },
          { id: 'Stage 05 Proposal', name: 'Stage 05 Proposal', icon: 'fa-file-signature', isStrategy: false }
        ];

        const PRIORITY_ORDER = { 'Urgent': 4, 'High': 3, 'Medium': 2, 'Low': 1 };
        const STATUS_ORDER = { 'New': 4, 'On progress': 3, 'Hold': 2, 'Completed': 1 };

        let state = {
          currentUser: null,
          customers: [],
          opportunities: [],
          tasks: [],
          audit_logs: [],
          selectedCustomerId: null,
          selectedOppId: null,
          activeStageIdx: 0,
          activeCommentTaskId: null,
          wizardIsNewCust: false,
          sortKey: 'priority',
          sortDirection: 'desc',
          calCurrentDate: new Date(),
          calSelectedDateStr: new Date().toISOString().split('T')[0]
        };

        function openModal(id) { document.getElementById(id).classList.remove('hidden'); }
        function closeModal(id) { document.getElementById(id).classList.add('hidden'); }

        function togglePasswordVisibility() {
          const input = document.getElementById('loginPassword');
          const icon = document.getElementById('passwordEyeIcon');
          if (input.type === 'password') {
            input.type = 'text';
            icon.className = 'fa-regular fa-eye-slash';
          } else {
            input.type = 'password';
            icon.className = 'fa-regular fa-eye';
          }
        }

        // --- AUTH & ROLES ---
        async function handleLogin(e) {
          e.preventDefault();
          const u = document.getElementById('loginUsername').value.trim();
          const p = document.getElementById('loginPassword').value.trim();
          const errDiv = document.getElementById('loginError');

          try {
            const res = await fetch('/api/auth/login', {
              method: 'POST',
              headers: {'Content-Type': 'application/json'},
              body: JSON.stringify({ username: u, password: p })
            });
            if (!res.ok) {
              const data = await res.json();
              errDiv.innerText = data.detail || 'Login failed';
              errDiv.classList.remove('hidden');
              return;
            }
            state.currentUser = await res.json();
            localStorage.setItem('presales_user', JSON.stringify(state.currentUser));
            document.getElementById('loginOverlay').classList.add('hidden');
            document.getElementById('appShell').classList.remove('hidden');
            applyUserRolePermissions();
            await fetchState();
          } catch (err) {
            errDiv.innerText = 'Server error during authentication.';
            errDiv.classList.remove('hidden');
          }
        }

        function handleLogout() {
          localStorage.removeItem('presales_user');
          state.currentUser = null;
          location.reload();
        }

        function applyUserRolePermissions() {
          if (!state.currentUser) return;
          const isAdmin = state.currentUser.role === 'admin';
          document.getElementById('roleBadge').innerText = state.currentUser.role.replace('_', ' ');
          document.getElementById('roleBadge').className = `text-[10px] font-extrabold px-2 py-0.5 rounded-full uppercase ${isAdmin ? 'bg-orange-50 text-orange-600 border border-orange-200' : 'bg-indigo-50 text-indigo-700 border border-indigo-200'}`;

          const adminOnlyIds = ['btnNewEngagement', 'btnUsersMgmt', 'btnEditCust', 'btnEditOpp'];
          adminOnlyIds.forEach(id => {
            const el = document.getElementById(id);
            if (el) {
              if (isAdmin) el.classList.remove('hidden'); else el.classList.add('hidden');
            }
          });
        }

        async function openUsersModal() {
          const res = await fetch('/api/users');
          const users = await res.json();
          const container = document.getElementById('usersListTable');
          container.innerHTML = users.map(u => `
            <div class="bg-slate-50 border border-slate-200 p-2.5 rounded-lg flex justify-between items-center">
              <div>
                <span class="font-bold text-slate-800">${u.full_name}</span> 
                <span class="text-slate-500">(@${u.username})</span>
              </div>
              <span class="px-2 py-0.5 rounded text-[10px] font-bold ${u.role === 'admin' ? 'bg-orange-100 text-orange-700' : 'bg-blue-100 text-blue-700'}">${u.role}</span>
            </div>
          `).join('');
          openModal('modalUsers');
        }

        async function handleCreateUser(e) {
          e.preventDefault();
          const username = document.getElementById('newUsername').value.trim();
          const password = document.getElementById('newPassword').value.trim();
          const full_name = document.getElementById('newFullName').value.trim();
          const role = document.getElementById('newRole').value;

          const res = await fetch('/api/users', {
            method: 'POST',
            headers: {'Content-Type': 'application/json'},
            body: JSON.stringify({ username, password, full_name, role })
          });
          if (res.ok) {
            alert('User created successfully!');
            document.getElementById('newUsername').value = '';
            document.getElementById('newPassword').value = '';
            document.getElementById('newFullName').value = '';
            await openUsersModal();
          } else {
            const d = await res.json();
            alert(d.detail || 'Failed to create user');
          }
        }

        // --- APP STATE & TREE ---
        async function fetchState() {
          const res = await fetch('/api/state');
          const data = await res.json();
          state.customers = data.customers;
          state.opportunities = data.opportunities;
          state.tasks = data.tasks;
          state.audit_logs = data.audit_logs;

          if (!state.selectedCustomerId && state.customers.length > 0) {
            state.selectedCustomerId = state.customers[0].id;
          }
          const validOpps = state.opportunities.filter(o => o.customer_id === state.selectedCustomerId);
          if (validOpps.length > 0 && !validOpps.some(o => o.id === state.selectedOppId)) {
            state.selectedOppId = validOpps[0].id;
          }

          renderExecutiveDashboard();
          renderNavigatorTree();
          renderStageWorkstream();
          renderAuditTrail();
        }

        function switchTab(tab) {
          ['Exec', 'Stage', 'Audit'].forEach(t => {
            document.getElementById('view' + t).classList.add('hidden');
            const btn = document.getElementById('tabBtn' + t);
            btn.className = 'px-4 py-2 rounded-xl text-xs font-bold bg-slate-100 hover:bg-slate-200 text-slate-700 border border-slate-200 hover:-translate-y-0.5 transition';
          });
          document.getElementById('view' + tab.charAt(0).toUpperCase() + tab.slice(1)).classList.remove('hidden');
          const activeBtn = document.getElementById('tabBtn' + tab.charAt(0).toUpperCase() + tab.slice(1));
          activeBtn.className = 'px-4 py-2 rounded-xl text-xs font-bold bg-gradient-to-r from-orange-400 to-amber-500 text-white shadow-md shadow-orange-400/20 hover:-translate-y-0.5 transition';
        }

        function openWizardModal(preselectedCustId = null) {
          if (state.currentUser?.role === 'view_only') {
            alert('View-only accounts cannot create new engagements.');
            return;
          }
          state.wizardIsNewCust = false;
          setWizardMode('existing');
          
          const sel = document.getElementById('wizCustSelect');
          sel.innerHTML = state.customers.map(c => 
            `<option value="${c.id}" ${c.id === (preselectedCustId || state.selectedCustomerId) ? 'selected' : ''}>${c.name}</option>`
          ).join('');

          document.getElementById('wizNewCustName').value = '';
          document.getElementById('wizNewCustIndustry').value = '';
          document.getElementById('wizOppTitle').value = '';
          document.getElementById('wizOppEnd').value = '';
          document.getElementById('wizOppPriority').value = 'Medium';
          document.getElementById('wizOppStatus').value = 'New';

          openModal('modalWizard');
        }

        function setWizardMode(mode) {
          state.wizardIsNewCust = mode === 'new';
          const btnEx = document.getElementById('wizBtnExisting');
          const btnNew = document.getElementById('wizBtnNew');
          const grpEx = document.getElementById('wizExistingGroup');
          const grpNew = document.getElementById('wizNewGroup');

          if (mode === 'existing') {
            btnEx.className = "py-2 text-xs font-bold rounded-lg bg-white text-orange-500 shadow-sm transition";
            btnNew.className = "py-2 text-xs font-bold rounded-lg text-slate-600 hover:text-slate-900 transition";
            grpEx.classList.remove('hidden');
            grpNew.classList.add('hidden');
          } else {
            btnNew.className = "py-2 text-xs font-bold rounded-lg bg-white text-orange-500 shadow-sm transition";
            btnEx.className = "py-2 text-xs font-bold rounded-lg text-slate-600 hover:text-slate-900 transition";
            grpNew.classList.remove('hidden');
            grpEx.classList.add('hidden');
          }
        }

        async function submitWizard() {
          const opp_title = document.getElementById('wizOppTitle').value.trim();
          const priority = document.getElementById('wizOppPriority').value;
          const status = document.getElementById('wizOppStatus').value;
          const end_date = document.getElementById('wizOppEnd').value;

          if (!opp_title) {
            alert('Please provide an opportunity title.');
            return;
          }

          let payload = {
            is_new_customer: state.wizardIsNewCust,
            opp_title,
            priority,
            status,
            end_date
          };

          if (state.wizardIsNewCust) {
            const customer_name = document.getElementById('wizNewCustName').value.trim();
            if (!customer_name) {
              alert('Please enter the customer name.');
              return;
            }
            payload.customer_name = customer_name;
            payload.industry = document.getElementById('wizNewCustIndustry').value.trim() || 'General';
          } else {
            payload.customer_id = parseInt(document.getElementById('wizCustSelect').value);
          }

          const res = await fetch('/api/engagements/unified', {
            method: 'POST',
            headers: {'Content-Type': 'application/json'},
            body: JSON.stringify(payload)
          });
          const data = await res.json();
          
          state.selectedCustomerId = data.customer_id;
          state.selectedOppId = data.opportunity_id;
          
          closeModal('modalWizard');
          await fetchState();
          switchTab('stage');
        }

        function renderNavigatorTree() {
          const search = (document.getElementById('navSearchInput')?.value || '').toLowerCase();
          const container = document.getElementById('navigatorTreeList');
          if (!container) return;

          container.innerHTML = '';

          state.customers.forEach(cust => {
            const opps = state.opportunities.filter(o => o.customer_id === cust.id);
            const matchesCust = cust.name.toLowerCase().includes(search);
            const matchingOpps = opps.filter(o => o.title.toLowerCase().includes(search));

            if (search && !matchesCust && matchingOpps.length === 0) return;

            const isCustActive = cust.id === state.selectedCustomerId;

            const card = document.createElement('div');
            card.className = `p-3.5 rounded-2xl border transition shadow-sm ${isCustActive ? 'bg-orange-50/40 border-orange-200' : 'bg-white border-slate-200/90 hover:bg-slate-50'}`;

            const initials = cust.name.split(' ').map(n => n[0]).slice(0, 2).join('').toUpperCase();
            const isAdmin = state.currentUser?.role === 'admin';

            card.innerHTML = `
              <div class="flex justify-between items-center pb-2.5">
                <div class="flex items-center gap-2 cursor-pointer" onclick="selectCustomerOnly(${cust.id})">
                  <div class="w-7 h-7 rounded-lg bg-gradient-to-tr from-orange-400 to-indigo-600 text-white flex items-center justify-center font-black text-[10px] shadow-sm">
                    ${initials}
                  </div>
                  <div>
                    <span class="font-extrabold text-xs text-slate-900 block line-clamp-1">${cust.name}</span>
                    <span class="text-[9px] text-slate-500 font-semibold">${cust.industry || 'Enterprise'}</span>
                  </div>
                </div>
                ${isAdmin ? `
                <div class="flex items-center gap-1">
                  <button onclick="openWizardModal(${cust.id})" class="text-[10px] text-orange-600 hover:text-white hover:bg-orange-500 font-bold px-2 py-0.5 rounded-lg border border-orange-200 transition" title="Add Deal">
                    + Deal
                  </button>
                  <button onclick="openEditCustomerModalById(${cust.id})" class="text-slate-400 hover:text-slate-600 p-1 text-xs">
                    <i class="fa-solid fa-pen"></i>
                  </button>
                </div>` : ''}
              </div>

              <div class="space-y-1.5 pl-2.5 border-l-2 border-slate-200 ml-2">
                ${opps.length === 0 ? `
                  <div class="text-[10px] text-slate-400 italic py-1">No deals yet. Click '+ Deal' above.</div>
                ` : opps.map(o => {
                  const isOppActive = o.id === state.selectedOppId;
                  return `
                    <div onclick="selectOpportunityDirect(${cust.id},${o.id})" class="p-2.5 rounded-xl cursor-pointer transition flex justify-between items-center ${
                      isOppActive ? 'bg-gradient-to-r from-orange-400 to-amber-500 text-white shadow-md shadow-orange-400/20' : 'bg-slate-50 hover:bg-slate-100 border border-slate-200 text-slate-700'
                    }">
                      <div class="truncate pr-2">
                        <div class="text-[11px] font-bold truncate">${o.title}</div>
                        <div class="text-[9px] ${isOppActive ? 'text-orange-100' : 'text-slate-400'} flex items-center gap-1.5 mt-0.5">
                          <i class="fa-regular fa-clock text-[8px]"></i> ${o.end_date || 'No Date'}
                        </div>
                      </div>
                      <span class="px-2 py-0.5 rounded text-[9px] font-black shrink-0 ${ isOppActive ? 'bg-white/20 text-white' : o.priority === 'Urgent' ? 'bg-rose-100 text-rose-700' : 'bg-orange-100 text-orange-700' }">${o.priority}</span>
                    </div>
                  `;
                }).join('')}
              </div>
            `;
            container.appendChild(card);
          });
        }

        function selectCustomerOnly(custId) {
          state.selectedCustomerId = custId;
          const opps = state.opportunities.filter(o => o.customer_id === custId);
          state.selectedOppId = opps.length > 0 ? opps[0].id : null;
          renderNavigatorTree();
          renderStageWorkstream();
        }

        function selectOpportunityDirect(custId, oppId) {
          state.selectedCustomerId = custId;
          state.selectedOppId = oppId;
          renderNavigatorTree();
          renderStageWorkstream();
        }

        function openEditCustomerModalById(custId) {
          state.selectedCustomerId = custId;
          openEditCustomerModal();
        }

        function renderStageWorkstream() {
          const opp = state.opportunities.find(o => o.id === state.selectedOppId);
          const cust = state.customers.find(c => c.id === state.selectedCustomerId);
          const isAdmin = state.currentUser?.role === 'admin';

          const bannerCustName = document.getElementById('bannerCustName');
          const bannerOppTitle = document.getElementById('bannerOppTitle');
          const bannerOppDate = document.getElementById('bannerOppDate');
          const bannerOppPriority = document.getElementById('bannerOppPriority');
          const bannerOppStatus = document.getElementById('bannerOppStatus');
          const area = document.getElementById('stageContentArea');
          const pillCont = document.getElementById('stagePillsContainer');

          if (!opp) {
            bannerCustName.innerText = cust ? cust.name : 'Account';
            bannerOppTitle.innerText = 'No Opportunity Selected';
            bannerOppDate.innerText = '-';
            bannerOppPriority.innerText = '-';
            bannerOppStatus.innerText = '-';
            pillCont.innerHTML = '';
            area.innerHTML = `
              <div class="text-center py-20 space-y-3">
                <div class="w-16 h-16 rounded-3xl bg-slate-100 text-orange-400 mx-auto flex items-center justify-center text-2xl">
                  <i class="fa-solid fa-folder-open"></i>
                </div>
                <div class="text-sm font-bold text-slate-800">No active workstream selected</div>
                <p class="text-xs text-slate-500">Choose a deal from the left navigator to manage deliverables.</p>
              </div>
            `;
            return;
          }

          bannerCustName.innerText = cust ? cust.name : 'Unknown Account';
          bannerOppTitle.innerText = opp.title;
          bannerOppDate.innerText = opp.end_date || 'Not scheduled';
          bannerOppPriority.innerText = opp.priority;
          bannerOppStatus.innerText = opp.status;

          pillCont.innerHTML = STAGES.map((s, idx) => {
            const isSelected = idx === state.activeStageIdx;
            const count = s.isStrategy ? '' : `<span class="ml-1.5 px-2 py-0.2 rounded-full text-[10px] font-black ${isSelected ? 'bg-white text-orange-600' : 'bg-slate-200 text-slate-700'}">${state.tasks.filter(t => t.opportunity_id === opp.id && t.stage_id === s.id).length}</span>`;
            return `
              <button onclick="setStageTab(${idx})" class="px-4 py-2 rounded-xl text-xs font-bold flex items-center gap-1.5 transition ${
                isSelected ? 'bg-gradient-to-r from-orange-400 to-amber-500 text-white shadow-md shadow-orange-400/20' : 'bg-slate-100 hover:bg-slate-200 text-slate-700 border border-slate-200'
              }">
                <i class="fa-solid ${s.icon} text-[11px] ${isSelected ? 'text-white' : 'text-slate-400'}"></i>
                <span>${s.name}</span>
                ${count}
              </button>
            `;
          }).join('');

          const activeStage = STAGES[state.activeStageIdx];

          if (activeStage.isStrategy) {
            area.innerHTML = `
              <div class="space-y-4">
                <div class="flex justify-between items-center pb-3 border-b border-slate-100">
                  <div class="flex items-center gap-2">
                    <div class="w-8 h-8 rounded-xl bg-orange-100 text-orange-600 flex items-center justify-center text-xs">
                      <i class="fa-solid fa-lightbulb"></i>
                    </div>
                    <div>
                      <h3 class="text-sm font-bold text-slate-900">Technical Win Strategy</h3>
                      <p class="text-xs text-slate-500">Positioning battlecard, competitor displacement, and solution goals.</p>
                    </div>
                  </div>
                  ${isAdmin ? `
                  <button onclick="saveStrategy()" class="bg-gradient-to-r from-orange-400 to-amber-500 hover:from-orange-500 hover:to-amber-600 text-white text-xs font-bold px-4 py-2 rounded-xl shadow-md transition">
                    <i class="fa-solid fa-floppy-disk mr-1"></i> Save Strategy
                  </button>` : ''}
                </div>
                <textarea id="strategyTextarea" rows="12" ${isAdmin ? '' : 'readonly'} placeholder="Write win strategy notes..." class="w-full text-xs p-4 rounded-2xl border border-slate-200 bg-slate-50 text-slate-900 focus:bg-white focus:ring-2 focus:ring-orange-400 outline-none leading-relaxed shadow-inner">${opp.strategy_notes || ''}</textarea>
              </div>
            `;
          } else {
            const tasks = state.tasks.filter(t => t.opportunity_id === opp.id && t.stage_id === activeStage.id);
            area.innerHTML = `
              <div class="space-y-4">
                <div class="flex justify-between items-center pb-3 border-b border-slate-100">
                  <div class="flex items-center gap-2">
                    <div class="w-8 h-8 rounded-xl bg-indigo-100 text-indigo-600 flex items-center justify-center text-xs">
                      <i class="fa-solid ${activeStage.icon}"></i>
                    </div>
                    <div>
                      <h3 class="text-sm font-bold text-slate-900">${activeStage.name} Workstream</h3>
                      <p class="text-xs text-slate-500">Milestone deliverables, target deadlines, and daily feedback.</p>
                    </div>
                  </div>
                  ${isAdmin ? `
                  <button onclick="openAddTaskModal()" class="bg-gradient-to-r from-orange-400 to-amber-500 hover:from-orange-500 hover:to-amber-600 text-white text-xs font-bold px-4 py-2 rounded-xl shadow-md transition">
                    <i class="fa-solid fa-plus mr-1"></i> Add Task
                  </button>` : ''}
                </div>
                <div class="overflow-x-auto">
                  <table class="w-full text-left text-xs">
                    <thead class="bg-slate-50 text-slate-500 uppercase text-[10px] tracking-wider border-b border-slate-200">
                      <tr>
                        <th class="py-2.5 px-3">Task Deliverable</th>
                        <th class="py-2.5 px-3">Start Date</th>
                        <th class="py-2.5 px-3">Target End</th>
                        <th class="py-2.5 px-3">Status</th>
                        <th class="py-2.5 px-3">Daily Updates</th>
                        ${isAdmin ? '<th class="py-2.5 px-3 text-right">Actions</th>' : ''}
                      </tr>
                    </thead>
                    <tbody class="divide-y divide-slate-100 font-medium">
                      ${tasks.length === 0 ? `
                        <tr>
                          <td colspan="6" class="text-center py-12 text-slate-400">
                            No deliverables logged for this stage yet. Click '+ Add Task' above.
                          </td>
                        </tr>
                      ` : tasks.map(t => `
                        <tr class="hover:bg-slate-50 transition">
                          <td class="py-3 px-3 font-bold text-slate-800">${t.title}</td>
                          <td class="py-3 px-3 text-slate-500">${t.start_date || '-'}</td>
                          <td class="py-3 px-3 text-slate-500">${t.end_date || '-'}</td>
                          <td class="py-3 px-3">
                            <span class="px-2.5 py-0.5 rounded text-[10px] font-black ${ t.status === 'Completed' ? 'bg-emerald-100 text-emerald-700' : t.status === 'On going' ? 'bg-blue-100 text-blue-700' : 'bg-orange-100 text-orange-700' }">${t.status}</span>
                          </td>
                          <td class="py-3 px-3">
                            <button onclick="openCommentModal(${t.id})" class="inline-flex items-center gap-1.5 px-2.5 py-1 bg-slate-100 hover:bg-slate-200 text-slate-700 rounded-lg font-bold text-[11px] transition">
                              <i class="fa-regular fa-comment-dots text-orange-500"></i> ${t.comments.length} Logged
                            </button>
                          </td>
                          ${isAdmin ? `
                          <td class="py-3 px-3 text-right space-x-1">
                            <button onclick="openEditTaskModal(${t.id})" class="p-1.5 hover:bg-slate-200 rounded-lg text-slate-500 hover:text-slate-800 transition"><i class="fa-solid fa-pen-to-square"></i></button>
                            <button onclick="deleteTask(${t.id})" class="p-1.5 hover:bg-rose-100 rounded-lg text-slate-500 hover:text-rose-600 transition"><i class="fa-solid fa-trash-can"></i></button>
                          </td>` : ''}
                        </tr>
                      `).join('')}
                    </tbody>
                  </table>
                </div>
              </div>
            `;
          }
        }

        function setStageTab(idx) {
          state.activeStageIdx = idx;
          renderStageWorkstream();
        }

        // --- EDIT MODALS ---
        function openEditCustomerModal() {
          if (state.currentUser?.role === 'view_only') return;
          const cust = state.customers.find(c => c.id === state.selectedCustomerId);
          if (!cust) return;
          document.getElementById('editCustId').value = cust.id;
          document.getElementById('editCustName').value = cust.name;
          document.getElementById('editCustIndustry').value = cust.industry || 'General';
          openModal('modalEditCustomer');
        }

        async function submitEditCustomer() {
          const id = document.getElementById('editCustId').value;
          const name = document.getElementById('editCustName').value.trim();
          const industry = document.getElementById('editCustIndustry').value.trim();
          if (!name) return;

          await fetch(`/api/customers/${id}`, {
            method: 'PUT',
            headers: {'Content-Type': 'application/json'},
            body: JSON.stringify({ name, industry })
          });
          closeModal('modalEditCustomer');
          await fetchState();
        }

        function openEditCurrentOppModal() {
          if (state.currentUser?.role === 'view_only') return;
          if (!state.selectedOppId) return;
          openEditOppModalById(state.selectedOppId);
        }

        function openEditOppModalById(oppId) {
          if (state.currentUser?.role === 'view_only') return;
          const opp = state.opportunities.find(o => o.id === oppId);
          if (!opp) return;
          document.getElementById('editOppId').value = opp.id;
          document.getElementById('editOppTitle').value = opp.title;
          document.getElementById('editOppEnd').value = opp.end_date || '';
          document.getElementById('editOppPriority').value = opp.priority;
          document.getElementById('editOppStatus').value = opp.status;
          openModal('modalEditOpp');
        }

        async function submitEditOpportunity() {
          const id = document.getElementById('editOppId').value;
          const title = document.getElementById('editOppTitle').value.trim();
          const end_date = document.getElementById('editOppEnd').value;
          const priority = document.getElementById('editOppPriority').value;
          const status = document.getElementById('editOppStatus').value;

          if (!title) return;

          await fetch(`/api/opportunities/${id}`, {
            method: 'PATCH',
            headers: {'Content-Type': 'application/json'},
            body: JSON.stringify({ title, end_date, priority, status })
          });

          closeModal('modalEditOpp');
          await fetchState();
        }

        // --- TASK HANDLERS ---
        function openAddTaskModal() {
          if (state.currentUser?.role === 'view_only') return;
          document.getElementById('taskEditId').value = '';
          document.getElementById('modalTaskTitle').innerText = 'Add Task Deliverable';
          document.getElementById('inputTaskTitle').value = '';
          document.getElementById('inputTaskStart').value = new Date().toISOString().split('T')[0];
          document.getElementById('inputTaskEnd').value = '';
          document.getElementById('inputTaskStatus').value = 'New';
          openModal('modalTask');
        }

        function openEditTaskModal(id) {
          if (state.currentUser?.role === 'view_only') return;
          const t = state.tasks.find(x => x.id === id);
          if (!t) return;
          document.getElementById('taskEditId').value = t.id;
          document.getElementById('modalTaskTitle').innerText = 'Edit Task Deliverable';
          document.getElementById('inputTaskTitle').value = t.title;
          document.getElementById('inputTaskStart').value = t.start_date || '';
          document.getElementById('inputTaskEnd').value = t.end_date || '';
          document.getElementById('inputTaskStatus').value = t.status;
          openModal('modalTask');
        }

        async function submitTaskModal() {
          const id = document.getElementById('taskEditId').value;
          const title = document.getElementById('inputTaskTitle').value.trim();
          const start_date = document.getElementById('inputTaskStart').value;
          const end_date = document.getElementById('inputTaskEnd').value;
          const status = document.getElementById('inputTaskStatus').value;
          if (!title) return;

          if (id) {
            await fetch(`/api/tasks/${id}`, {
              method: 'PUT',
              headers: {'Content-Type': 'application/json'},
              body: JSON.stringify({ title, start_date, end_date, status })
            });
          } else {
            await fetch('/api/tasks', {
              method: 'POST',
              headers: {'Content-Type': 'application/json'},
              body: JSON.stringify({
                opportunity_id: state.selectedOppId,
                stage_id: STAGES[state.activeStageIdx].id,
                title, start_date, end_date, status
              })
            });
          }
          closeModal('modalTask');
          await fetchState();
        }

        async function deleteTask(id) {
          if (state.currentUser?.role === 'view_only') return;
          if (confirm('Delete this deliverable?')) {
            await fetch(`/api/tasks/${id}`, { method: 'DELETE' });
            await fetchState();
          }
        }

        function openCommentModal(id) {
          state.activeCommentTaskId = id;
          const t = state.tasks.find(x => x.id === id);
          document.getElementById('commentModalTaskTitle').innerText = t ? t.title : '';
          const list = document.getElementById('commentList');
          list.innerHTML = (t && t.comments.length > 0) ? t.comments.map(c => `
            <div class="bg-slate-50 border border-slate-200 p-3 rounded-xl space-y-1">
              <span class="text-[10px] font-bold text-orange-500 flex items-center gap-1"><i class="fa-regular fa-clock text-[9px]"></i> ${c.date}</span>
              <p class="text-slate-800 font-medium">${c.text}</p>
            </div>
          `).join('') : '<div class="text-center py-6 text-slate-400">No updates logged yet.</div>';
          
          const isAdmin = state.currentUser?.role === 'admin';
          document.getElementById('commentFormContainer').style.display = isAdmin ? 'block' : 'none';
          openModal('modalComment');
        }

        async function submitTaskComment() {
          const notes = document.getElementById('commentInputText').value.trim();
          if (!notes) return;
          await fetch(`/api/tasks/${state.activeCommentTaskId}/comments`, {
            method: 'POST',
            headers: {'Content-Type': 'application/json'},
            body: JSON.stringify({ notes })
          });
          document.getElementById('commentInputText').value = '';
          closeModal('modalComment');
          await fetchState();
        }

        // --- DRILLDOWN WINDOWS ---
        function openDrilldownModal(type) {
          const titleEl = document.getElementById('tileModalTitle');
          const subEl = document.getElementById('tileModalSubtitle');
          const content = document.getElementById('tileModalContent');
          content.innerHTML = '';

          if (type === 'engaged') {
            titleEl.innerText = 'Active Engaged Accounts';
            subEl.innerText = 'Customers with active deals in motion';
            
            const activeCusts = state.customers.filter(c => 
              state.opportunities.some(o => o.customer_id === c.id && o.status !== 'Completed')
            );

            if (activeCusts.length === 0) {
              content.innerHTML = `<div class="text-center py-8 text-slate-400">No active customer engagements found.</div>`;
            } else {
              content.innerHTML = activeCusts.map(c => {
                const opps = state.opportunities.filter(o => o.customer_id === c.id && o.status !== 'Completed');
                return `
                  <div class="bg-slate-50 border-2 border-slate-200 p-3.5 rounded-2xl flex justify-between items-center shadow-sm">
                    <div>
                      <div class="font-bold text-slate-900 text-sm">${c.name}</div>
                      <div class="text-xs text-slate-500">${c.industry || 'General'} &bull; <span class="font-bold text-orange-500">${opps.length} Active Deal${opps.length > 1 ? 's' : ''}</span></div>
                    </div>
                    <button onclick="closeModal('modalTileDetails'); jumpToOpportunity(${c.id}, ${opps[0]?.id || ''})" class="px-3.5 py-1.5 bg-orange-500 hover:bg-orange-600 text-white font-bold rounded-xl text-xs transition shadow-sm">
                      Workstream &rarr;
                    </button>
                  </div>
                `;
              }).join('');
            }
          } 
          else if (type === 'active_opps') {
            titleEl.innerText = 'Active Pipeline Deals';
            subEl.innerText = 'All ongoing customer deals across stages';

            const activeOpps = state.opportunities.filter(o => o.status !== 'Completed');
            if (activeOpps.length === 0) {
              content.innerHTML = `<div class="text-center py-8 text-slate-400">No active deals found.</div>`;
            } else {
              content.innerHTML = activeOpps.map(o => {
                const cust = state.customers.find(c => c.id === o.customer_id);
                return `
                  <div class="bg-slate-50 border-2 border-slate-200 p-3.5 rounded-2xl flex justify-between items-center shadow-sm">
                    <div>
                      <div class="font-bold text-slate-900 text-sm">${o.title}</div>
                      <div class="text-xs text-slate-500">${cust ? cust.name : 'Unknown'} &bull; Target: <span class="font-bold text-slate-800">${o.end_date || 'No Date'}</span></div>
                    </div>
                    <div class="flex items-center gap-2">
                      <span class="px-2.5 py-0.5 rounded text-[10px] font-bold ${o.priority === 'Urgent' ? 'bg-rose-100 text-rose-700' : 'bg-orange-100 text-orange-700'}">${o.priority}</span>
                      <button onclick="closeModal('modalTileDetails'); jumpToOpportunity(${o.customer_id}, ${o.id})" class="px-3.5 py-1.5 bg-slate-800 hover:bg-slate-900 text-white font-bold rounded-xl text-xs transition shadow-sm">
                        Open &rarr;
                      </button>
                    </div>
                  </div>
                `;
              }).join('');
            }
          }
          else if (type === 'urgent') {
            titleEl.innerText = 'High & Urgent Priority Deals';
            subEl.innerText = 'Critical focus deals requiring attention';

            const urgentOpps = state.opportunities.filter(o => ['Urgent', 'High'].includes(o.priority) && o.status !== 'Completed');
            if (urgentOpps.length === 0) {
              content.innerHTML = `<div class="text-center py-8 text-slate-400">No urgent deals found.</div>`;
            } else {
              content.innerHTML = urgentOpps.map(o => {
                const cust = state.customers.find(c => c.id === o.customer_id);
                return `
                  <div class="bg-rose-50 border-2 border-rose-200 p-3.5 rounded-2xl flex justify-between items-center shadow-sm">
                    <div>
                      <div class="font-bold text-slate-900 text-sm">${o.title}</div>
                      <div class="text-xs text-slate-500">${cust ? cust.name : 'Unknown'} &bull; Target: <span class="font-bold text-rose-700">${o.end_date || 'No Date'}</span></div>
                    </div>
                    <button onclick="closeModal('modalTileDetails'); jumpToOpportunity(${o.customer_id}, ${o.id})" class="px-3.5 py-1.5 bg-rose-600 hover:bg-rose-700 text-white font-bold rounded-xl text-xs transition shadow-sm">
                      Open &rarr;
                    </button>
                  </div>
                `;
              }).join('');
            }
          }
          else if (type === 'completed') {
            titleEl.innerText = 'Delivered & Completed Tasks';
            subEl.innerText = 'Milestones completed across client workstreams';

            const completed = state.tasks.filter(t => t.status === 'Completed');
            if (completed.length === 0) {
              content.innerHTML = `<div class="text-center py-8 text-slate-400">No completed tasks yet.</div>`;
            } else {
              content.innerHTML = completed.map(t => {
                const opp = state.opportunities.find(o => o.id === t.opportunity_id);
                const cust = opp ? state.customers.find(c => c.id === opp.customer_id) : null;
                return `
                  <div class="bg-slate-50 border-2 border-slate-200 p-3.5 rounded-2xl flex justify-between items-center shadow-sm">
                    <div>
                      <div class="font-bold text-slate-900 text-xs">${t.title}</div>
                      <div class="text-[11px] text-slate-500">${cust ? cust.name : 'Unknown'} &bull; ${opp ? opp.title : ''} &bull; <span class="font-bold text-orange-500">${t.stage_id}</span></div>
                    </div>
                    <span class="px-2.5 py-0.5 rounded text-[10px] font-bold bg-emerald-100 text-emerald-700 border border-emerald-200">Completed</span>
                  </div>
                `;
              }).join('');
            }
          }

          openModal('modalTileDetails');
        }

        function openStatusTasksModal(status) {
          const titleEl = document.getElementById('tileModalTitle');
          const subEl = document.getElementById('tileModalSubtitle');
          const content = document.getElementById('tileModalContent');
          
          titleEl.innerText = `Tasks: ${status}`;
          subEl.innerText = `All tasks currently marked '${status}'`;

          const filtered = state.tasks.filter(t => t.status === status);
          if (filtered.length === 0) {
            content.innerHTML = `<div class="text-center py-8 text-slate-400">No tasks marked as '${status}'.</div>`;
          } else {
            content.innerHTML = filtered.map(t => {
              const opp = state.opportunities.find(o => o.id === t.opportunity_id);
              const cust = opp ? state.customers.find(c => c.id === opp.customer_id) : null;
              return `
                <div class="bg-slate-50 border-2 border-slate-200 p-3 rounded-xl flex justify-between items-center shadow-sm">
                  <div>
                    <div class="font-bold text-slate-900 text-xs">${t.title}</div>
                    <div class="text-[11px] text-slate-500">${cust ? cust.name : 'Unknown'} &bull; ${opp ? opp.title : ''} &bull; <span class="font-bold text-orange-500">${t.stage_id}</span></div>
                  </div>
                  <button onclick="closeModal('modalTileDetails'); jumpToOpportunity(${cust ? cust.id : ''}, ${opp ? opp.id : ''})" class="px-3 py-1 bg-slate-800 hover:bg-slate-900 text-white rounded-lg text-xs font-semibold shadow-sm">
                    View &rarr;
                  </button>
                </div>
              `;
            }).join('');
          }
          openModal('modalTileDetails');
        }

        function openStageTasksModal(stageId) {
          const titleEl = document.getElementById('tileModalTitle');
          const subEl = document.getElementById('tileModalSubtitle');
          const content = document.getElementById('tileModalContent');
          
          titleEl.innerText = `${stageId} Tasks`;
          subEl.innerText = `All deliverables currently active within ${stageId}`;

          const filtered = state.tasks.filter(t => t.stage_id === stageId);
          if (filtered.length === 0) {
            content.innerHTML = `<div class="text-center py-8 text-slate-400">No tasks under ${stageId}.</div>`;
          } else {
            content.innerHTML = filtered.map(t => {
              const opp = state.opportunities.find(o => o.id === t.opportunity_id);
              const cust = opp ? state.customers.find(c => c.id === opp.customer_id) : null;
              return `
                <div class="bg-slate-50 border-2 border-slate-200 p-3 rounded-xl flex justify-between items-center shadow-sm">
                  <div>
                    <div class="font-bold text-slate-900 text-xs">${t.title}</div>
                    <div class="text-[11px] text-slate-500">${cust ? cust.name : 'Unknown'} &bull; ${opp ? opp.title : ''}</div>
                  </div>
                  <button onclick="closeModal('modalTileDetails'); jumpToOpportunity(${cust ? cust.id : ''}, ${opp ? opp.id : ''})" class="px-3 py-1 bg-orange-500 hover:bg-orange-600 text-white rounded-lg text-xs font-semibold shadow-sm">
                    View &rarr;
                  </button>
                </div>
              `;
            }).join('');
          }
          openModal('modalTileDetails');
        }

        function renderExecutiveDashboard() {
          const activeCusts = new Set(state.opportunities.filter(o => o.status !== 'Completed').map(o => o.customer_id)).size;
          const activeOpps = state.opportunities.filter(o => o.status !== 'Completed').length;
          const urgentOpps = state.opportunities.filter(o => ['Urgent', 'High'].includes(o.priority) && o.status !== 'Completed').length;
          const completedTasks = state.tasks.filter(t => t.status === 'Completed').length;

          document.getElementById('kpiCustCount').innerText = activeCusts;
          document.getElementById('kpiOppCount').innerText = activeOpps;
          document.getElementById('kpiUrgentCount').innerText = urgentOpps;
          document.getElementById('kpiCompletedTasks').innerText = completedTasks;

          const statuses = ['New', 'On going', 'Hold', 'Completed'];
          const statusCont = document.getElementById('statusTilesContainer');
          statusCont.innerHTML = statuses.map(s => {
            const count = state.tasks.filter(t => t.status === s).length;
            return `
              <button onclick="openStatusTasksModal('${s}')" class="p-3.5 rounded-xl border-2 border-slate-200 bg-white hover:border-orange-400 text-left transition shadow-sm hover:shadow hover:-translate-y-0.5 group">
                <span class="text-[10px] uppercase font-extrabold text-slate-500 group-hover:text-orange-500 block">${s}</span>
                <span class="text-xl font-black text-slate-900 mt-1 block">${count}</span>
              </button>
            `;
          }).join('');

          const stageCont = document.getElementById('stageTilesContainer');
          stageCont.innerHTML = STAGES.slice(1).map(s => {
            const count = state.tasks.filter(t => t.stage_id === s.id).length;
            return `
              <button onclick="openStageTasksModal('${s.id}')" class="p-3.5 rounded-xl border-2 border-slate-200 bg-white hover:border-indigo-400 text-left transition shadow-sm hover:shadow hover:-translate-y-0.5 group">
                <span class="text-[10px] uppercase font-extrabold text-slate-500 group-hover:text-indigo-600 block line-clamp-1">${s.name.replace(/Stage 0\d /, '')}</span>
                <span class="text-xl font-black text-slate-900 mt-1 block">${count}</span>
              </button>
            `;
          }).join('');

          updateSortIcons();

          let oppsToRender = [...state.opportunities].map(o => {
            const cust = state.customers.find(c => c.id === o.customer_id);
            return { ...o, customer_name: cust ? cust.name : 'Unknown' };
          });

          oppsToRender.sort((a, b) => {
            let res = 0;
            if (state.sortKey === 'priority') {
              res = (PRIORITY_ORDER[a.priority] || 0) - (PRIORITY_ORDER[b.priority] || 0);
            } else if (state.sortKey === 'status') {
              res = (STATUS_ORDER[a.status] || 0) - (STATUS_ORDER[b.status] || 0);
            } else if (state.sortKey === 'customer') {
              res = a.customer_name.localeCompare(b.customer_name);
            }
            return state.sortDirection === 'asc' ? res : -res;
          });

          const tbody = document.getElementById('execMatrixBody');
          tbody.innerHTML = oppsToRender.map(o => {
            const timeline = o.end_date ? 
              `<div class="font-bold text-slate-800 flex items-center gap-1.5"><i class="fa-regular fa-calendar-check text-orange-500"></i> ${o.end_date}</div>` : 
              `<span class="text-slate-400">Not set</span>`;

            const isAdmin = state.currentUser?.role === 'admin';

            return `
              <tr class="hover:bg-slate-50/80 transition">
                <td class="py-3 px-4 font-bold text-slate-900 cursor-pointer" onclick="jumpToOpportunity(${o.customer_id}, ${o.id})">
                  <div class="flex items-center gap-2">
                    <span class="w-2.5 h-2.5 rounded-full ${o.priority === 'Urgent' ? 'bg-rose-500' : 'bg-orange-400'}"></span>
                    <span>${o.customer_name}</span>
                  </div>
                </td>
                <td class="py-3 px-4 font-semibold text-slate-700 cursor-pointer" onclick="jumpToOpportunity(${o.customer_id}, ${o.id})">${o.title}</td>
                <td class="py-3 px-4">${timeline}</td>
                <td class="py-3 px-4">
                  <span class="px-2.5 py-0.5 rounded text-[10px] font-black ${
                    o.priority === 'Urgent' ? 'bg-rose-100 text-rose-700 border border-rose-200' : 
                    o.priority === 'High' ? 'bg-orange-100 text-orange-700 border border-orange-200' : 'bg-slate-100 text-slate-700'
                  }">${o.priority}</span>
                </td>
                <td class="py-3 px-4">
                  <span class="px-2.5 py-0.5 rounded text-[10px] font-bold ${
                    o.status === 'Completed' ? 'bg-emerald-100 text-emerald-700 border border-emerald-200' : 'bg-blue-100 text-blue-700 border border-blue-200'
                  }">${o.status}</span>
                </td>
                <td class="py-3 px-4 text-right space-x-2">
                  ${isAdmin ? `
                  <button onclick="openEditOppModalById(${o.id})" class="text-orange-600 hover:text-orange-800 font-bold text-xs px-2.5 py-1 bg-orange-50 hover:bg-orange-100 rounded-lg border border-orange-200 transition shadow-sm">
                    <i class="fa-solid fa-pen mr-1"></i> Edit
                  </button>` : ''}
                  <button onclick="jumpToOpportunity(${o.customer_id}, ${o.id})" class="text-indigo-600 hover:text-indigo-800 font-bold text-xs">
                    Engine &rarr;
                  </button>
                </td>
              </tr>
            `;
          }).join('');

          renderCalendar();
        }

        function updateSortIcons() {
          ['customer', 'priority', 'status'].forEach(k => {
            const el = document.getElementById('sortIcon' + k.charAt(0).toUpperCase() + k.slice(1));
            if (state.sortKey === k) {
              el.className = `fa-solid fa-sort-${state.sortDirection === 'asc' ? 'up' : 'down'} ml-1 text-orange-500`;
            } else {
              el.className = 'fa-solid fa-sort ml-1 text-slate-400';
            }
          });
          const badge = document.getElementById('sortStatusBadge');
          const dirText = state.sortDirection === 'asc' ? 'Low &rarr; High' : 'High &rarr; Low';
          badge.innerHTML = `Sorted by ${state.sortKey.toUpperCase()}: ${dirText}`;
        }

        function toggleSort(key) {
          if (state.sortKey === key) {
            state.sortDirection = state.sortDirection === 'asc' ? 'desc' : 'asc';
          } else {
            state.sortKey = key;
            state.sortDirection = (key === 'priority' || key === 'status') ? 'desc' : 'asc';
          }
          renderExecutiveDashboard();
        }

        function jumpToOpportunity(custId, oppId) {
          state.selectedCustomerId = custId;
          state.selectedOppId = oppId;
          switchTab('stage');
          renderNavigatorTree();
          renderStageWorkstream();
        }

        function changeCalMonth(delta) {
          state.calCurrentDate.setMonth(state.calCurrentDate.getMonth() + delta);
          renderCalendar();
        }

        function renderCalendar() {
          const year = state.calCurrentDate.getFullYear();
          const month = state.calCurrentDate.getMonth();
          const monthNames = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"];
          document.getElementById('calMonthYearLabel').innerText = `${monthNames[month]} ${year}`;

          const firstDay = new Date(year, month, 1).getDay();
          const daysInMonth = new Date(year, month + 1, 0).getDate();
          const grid = document.getElementById('calendarDaysGrid');
          grid.innerHTML = '';

          for (let i = 0; i < firstDay; i++) {
            grid.innerHTML += `<div class="h-16 bg-slate-100/50 rounded-xl border border-transparent"></div>`;
          }

          for (let d = 1; d <= daysInMonth; d++) {
            const dateStr = `${year}-${String(month + 1).padStart(2, '0')}-${String(d).padStart(2, '0')}`;
            const dayTasks = state.tasks.filter(t => {
              if (t.start_date === dateStr || t.end_date === dateStr) return true;
              if (t.start_date && t.end_date && dateStr >= t.start_date && dateStr <= t.end_date) return true;
              return false;
            });

            const isSelected = state.calSelectedDateStr === dateStr;
            const hasTasks = dayTasks.length > 0;

            grid.innerHTML += `
              <div onclick="selectCalDate('${dateStr}')" class="h-16 p-2 rounded-xl border-2 text-left cursor-pointer transition flex flex-col justify-between shadow-sm ${
                isSelected ? 'bg-orange-50 border-orange-400 ring-2 ring-orange-300' : 
                hasTasks ? 'bg-white hover:bg-slate-50 border-slate-200' : 'bg-white/80 hover:bg-slate-50 border-slate-200/60'
              }">
                <span class="text-[11px] font-black ${isSelected ? 'text-orange-600' : 'text-slate-700'}">${d}</span>
                ${hasTasks ? `
                  <div class="flex items-center gap-1">
                    <span class="w-2 h-2 rounded-full bg-orange-400"></span>
                    <span class="text-[9px] font-bold text-slate-600">${dayTasks.length} task${dayTasks.length > 1 ? 's' : ''}</span>
                  </div>
                ` : ''}
              </div>
            `;
          }

          renderCalendarAgenda();
        }

        function selectCalDate(dateStr) {
          state.calSelectedDateStr = dateStr;
          renderCalendar();
        }

        function renderCalendarAgenda() {
          const dateStr = state.calSelectedDateStr;
          document.getElementById('agendaSelectedDateLabel').innerText = `Tasks on ${dateStr}`;

          const dayTasks = state.tasks.filter(t => {
            if (t.start_date === dateStr || t.end_date === dateStr) return true;
            if (t.start_date && t.end_date && dateStr >= t.start_date && dateStr <= t.end_date) return true;
            return false;
          });

          document.getElementById('agendaTaskCountBadge').innerText = `${dayTasks.length} Task${dayTasks.length === 1 ? '' : 's'}`;
          const list = document.getElementById('agendaTaskList');

          if (dayTasks.length === 0) {
            list.innerHTML = `<div class="text-center py-12 text-slate-400 text-xs">No tasks scheduled for this day.</div>`;
            return;
          }

          list.innerHTML = dayTasks.map(t => {
            const opp = state.opportunities.find(o => o.id === t.opportunity_id);
            const cust = opp ? state.customers.find(c => c.id === opp.customer_id) : null;
            return `
              <div class="bg-white p-3 rounded-xl border-2 border-slate-200 space-y-1 shadow-sm">
                <div class="flex justify-between items-start">
                  <span class="text-[10px] font-extrabold text-indigo-700 uppercase">${cust ? cust.name : ''}</span>
                  <span class="px-2 py-0.2 rounded text-[9px] font-bold ${t.status === 'Completed' ? 'bg-emerald-100 text-emerald-700' : 'bg-blue-100 text-blue-700'}">${t.status}</span>
                </div>
                <div class="font-bold text-slate-800 text-xs">${t.title}</div>
                <div class="text-[10px] text-slate-500 flex justify-between items-center pt-1 border-t border-slate-100">
                  <span>${t.stage_id}</span>
                  <button onclick="jumpToOpportunity(${cust ? cust.id : ''}, ${opp ? opp.id : ''})" class="text-orange-600 font-bold hover:underline">View &rarr;</button>
                </div>
              </div>
            `;
          }).join('');
        }

        function renderAuditTrail() {
          const tbody = document.getElementById('auditTableBody');
          tbody.innerHTML = state.audit_logs.map(l => `
            <tr class="hover:bg-slate-50 transition">
              <td class="py-2.5 px-4 font-mono text-slate-400 text-[10px]">${l.created_at}</td>
              <td class="py-2.5 px-4 font-bold text-orange-500">${l.action_type}</td>
              <td class="py-2.5 px-4 text-slate-700 font-semibold">${l.customer_name || '-'}</td>
              <td class="py-2.5 px-4 text-slate-700 font-semibold">${l.opp_title || '-'}</td>
              <td class="py-2.5 px-4 text-slate-600">${l.details}</td>
            </tr>
          `).join('');
        }

        // --- RESUME SESSION ON RELOAD ---
        window.onload = async () => {
          const savedUser = localStorage.getItem('presales_user');
          if (savedUser) {
            try {
              state.currentUser = JSON.parse(savedUser);
              document.getElementById('loginOverlay').classList.add('hidden');
              document.getElementById('appShell').classList.remove('hidden');
              applyUserRolePermissions();
              await fetchState();
              return;
            } catch (e) {
              localStorage.removeItem('presales_user');
            }
          }
          document.getElementById('loginOverlay').classList.remove('hidden');
        };
      </script>
    </body>
    </html>
    """